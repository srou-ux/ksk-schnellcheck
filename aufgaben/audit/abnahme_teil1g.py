"""Abnahmetest für Chatbot Teil 1g (Mischform: feste Textbausteine). Vor dem Bau geschrieben: muss VORHER rot sein,
NACHHER 0 Lücken. Erfundene Daten. Aufruf: python aufgaben/audit/abnahme_teil1g.py"""
import sys, pathlib, re
W = pathlib.Path(__file__).resolve().parents[2]
sys.path.insert(0, str(W))
from decimal import Decimal as D
import openpyxl
from kern import rechnen as K
from tests.app.hilfen import lauf_bestaetigt
from app import router
from app.llm.fake import FakeLLM
from app.schema import Fakten, Position as P

NP, GMBH = "natuerliche_person", "GmbH_UG_AG"
LUECKEN = []


def eur(x):
    return f"{D(x):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def pruefe(name, bedingung, detail=""):
    print(f"{name:66} {'ok' if bedingung else 'LÜCKE'}" + ("" if bedingung else f"   <- {str(detail)[:160]}"))
    if not bedingung:
        LUECKEN.append(name)


class KeinKoerper(FakeLLM):
    """Der Antwort-Körper darf NICHT vom Modell kommen: formuliere_antwort ruft man nicht mehr auf.
    beschreibe_fall liefert absichtlich verbotenen Text (Ziffern, Paragraf, Ergebnisaussage, Link)."""
    def __init__(self, beschreibung):
        super().__init__(); self.beschreibung = beschreibung; self.beschreibe_aufrufe = 0; self.koerper_aufrufe = 0
    def formuliere_antwort(self, kontext):
        self.koerper_aufrufe += 1
        raise AssertionError("formuliere_antwort darf nicht mehr aufgerufen werden")
    def beschreibe_fall(self, kontext):
        self.beschreibe_aufrufe += 1
        return self.beschreibung


def lauf(text, fakten, beschreibung="Es geht um Aufträge an Freiberufler."):
    llm = KeinKoerper(beschreibung)
    try:
        b = lauf_bestaetigt(llm, text, fakten)
    except AssertionError as e:
        class B: antwort = str(e); ergebnisklasse = "-"; fundstellen = []
        return B(), llm
    return b, llm


# ---- G1: Blatt im Regelkatalog --------------------------------------------------------------
wb = openpyxl.load_workbook(W / "regelwerk" / "Regelkatalog_24_KSVG_Entwurf.xlsx", data_only=True)
pruefe("G1 Blatt 'Textbausteine' im Regelkatalog vorhanden", "Textbausteine" in wb.sheetnames)
texte = []
if "Textbausteine" in wb.sheetnames:
    ws = wb["Textbausteine"]
    zeilen = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]
    texte = [str(r[5]) for r in zeilen if r[5]]
    pruefe("G1b jede Zeile hat ID, Text und Prüfstatus", all(r[0] and r[5] and r[8] for r in zeilen), zeilen[:1])

# ---- G2: kein Festtext im Programm ------------------------------------------------------------
code = " ".join(p.read_text(encoding="utf-8") for p in (W / "app").glob("**/*.py"))
fest = [t for t in texte if len(t) > 25 and "{" in t and t in code]
pruefe("G2 Bausteintexte stehen nicht mehr im Programmcode (nur im Regelkatalog)", texte and not fest, fest[:2])

# ---- G3: Körper kommt ohne Modell, Zahlen aus dem Kern ---------------------------------------
t = "G3 gemischte Positionen 2025"
f = Fakten(anfrageart="bewertung", jahr=2025, verwerterart="typischer_verwerter", leistungsart="Grafik/Design",
           positionen=[P(betrag=D("2400"), rolle="honorar", auftragnehmer_typ=NP), P(betrag=D("1100"), rolle="honorar", auftragnehmer_typ=NP),
                       P(betrag=D("5000"), rolle="honorar", auftragnehmer_typ=GMBH, bezeichnung="GmbH-Studio"),
                       P(betrag=D("800"), rolle="honorar", auftragnehmer_typ=NP)])
bg = D("4300"); ab = K.berechne_abgabe(bg, 2025)[0]
b, llm = lauf(t, f)
pruefe("G3a Antwort wird ausgegeben (nicht 'Pruefung fehlgeschlagen')", "fehlgeschlagen" not in b.ergebnisklasse.lower() and "Abgabepflichtig" in b.ergebnisklasse, b.ergebnisklasse + " " + b.antwort)
pruefe("G3b Bemessungsgrundlage und Abgabe aus dem Kern stehen drin", eur(bg) in b.antwort and eur(ab) in b.antwort, b.antwort)
pruefe("G3c formuliere_antwort wurde nie aufgerufen", llm.koerper_aufrufe == 0, llm.koerper_aufrufe)
pruefe("G3d Fundstelle und Schlusssatz vorhanden", bool(b.fundstellen) and "Entwurf – vom Steuerberater zu prüfen" in b.antwort, b.antwort[-200:])

# ---- G4: Modelltext nur als Fallbeschreibung, verbotenes fliegt still raus -------------------
bose = "Das ist ganz sicher nicht abgabepflichtig, siehe § 54 und BSG 12 R 3/20, zahlen Sie 99 Euro, mehr unter http://ksk.example."
b, llm = lauf("G4 böse Beschreibung", f, bose)
pruefe("G4a beschreibe_fall wurde aufgerufen", llm.beschreibe_aufrufe >= 1, llm.beschreibe_aufrufe)
pruefe("G4b verbotener Modelltext erscheint nicht", all(x not in b.antwort for x in ("ganz sicher", "12 R 3/20", "99 Euro", "http")), b.antwort)
pruefe("G4c Antwort bleibt trotzdem vollständig und richtig", "fehlgeschlagen" not in b.ergebnisklasse.lower() and eur(ab) in b.antwort, b.antwort)
gut = "Es geht um vier Zahlungen an Gestalter im Jahr {JAHR}."
b, llm = lauf("G4 gute Beschreibung", f, "Es geht um Zahlungen an Gestalter im Jahr {JAHR}.")
pruefe("G4d harmlose Fallbeschreibung mit Platzhalter wird übernommen", "Zahlungen an Gestalter im Jahr 2025" in b.antwort, b.antwort)
b, llm = lauf("G4 Ergebnisaussage", f, "Damit ist die Sache nicht abgabepflichtig.")
pruefe("G4e Ergebnisaussage im Modelltext fliegt raus", "Damit ist die Sache" not in b.antwort and eur(ab) in b.antwort, b.antwort)

# ---- G5: Säumnis, Auskunft, Vorauszahlung, Fahrtkosten ----------------------------------------
b, _ = lauf("G5 Säumnis", Fakten(anfrageart="saeumniszuschlag", jahr=2025, rueckstand=D("2000"), saeumnis_monate=4))
pruefe("G5a Säumniszuschlag-Wert aus dem Kern", eur(K.saeumniszuschlag([D("2000")], 4, 2025).abgabe) in b.antwort and "Abgabesatz" not in b.antwort, b.antwort)
b, _ = lauf("Wie hoch ist der Abgabesatz im Jahr 2024?", Fakten(anfrageart="auskunft", jahr=2024, auskunft_thema="abgabesatz"))
pruefe("G5b Auskunft Abgabesatz 2024 kommt durch", "Auskunft" in b.ergebnisklasse and "5" in b.antwort, b.ergebnisklasse + b.antwort)
b, _ = lauf("G5 Fahrtkosten", Fakten(anfrageart="bewertung", jahr=2024, verwerterart="typischer_verwerter", auftragnehmer_typ=NP, leistungsart="Fotografie",
            positionen=[P(betrag=D("1500"), rolle="honorar"), P(betrag=D("300"), rolle="fahrtkosten_nachgewiesen", nachgewiesen=True)]))
pruefe("G5c Fahrtkosten: 1.500,00 € Bemessungsgrundlage, nicht 1.800,00 €", eur(1500) in b.antwort and eur(1800) not in b.antwort.replace("1.500,00", ""), b.antwort)

# ---- G6: Ausfall des Modells bei der Beschreibung -------------------------------------------
class Defekt(KeinKoerper):
    def beschreibe_fall(self, kontext):
        from app.llm.base import LLMFehlgeschlagen
        raise LLMFehlgeschlagen("Zeitüberschreitung")
d = Defekt("")
try:
    b = lauf_bestaetigt(d, "G6 Modell fällt aus", f)
    pruefe("G6 Modellausfall bei der Beschreibung: Antwort trotzdem komplett", eur(ab) in b.antwort and "fehlgeschlagen" not in b.ergebnisklasse.lower(), b.antwort)
except Exception as e:
    pruefe("G6 Modellausfall bei der Beschreibung: Antwort trotzdem komplett", False, repr(e))

# ---- G7: Schnittstelle ----------------------------------------------------------------------
from app.llm import base
pruefe("G7a LLM-Schnittstelle kennt beschreibe_fall", "beschreibe_fall" in open(base.__file__, encoding="utf-8").read())
mist = (W / "app" / "llm" / "mistral.py").read_text(encoding="utf-8")
pruefe("G7b Mistral-Adapter hat beschreibe_fall", "def beschreibe_fall" in mist)

print(f"\nLücken: {len(LUECKEN)}")
sys.exit(1 if LUECKEN else 0)
