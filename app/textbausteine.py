"""Liest das Blatt "Textbausteine" aus `regelwerk/Regelkatalog_24_KSVG_Entwurf.xlsx`
(Chatbot Teil 1g, Mischform). Rechtsaussagen/Koerpertexte stehen NUR hier, nie im
Programmcode -- das Sprachmodell schreibt nur noch eine kurze Fallbeschreibung
(`app/llm/*::beschreibe_fall`), nie den eigentlichen Antwort-Koerper mehr
(`formuliere_antwort` wird nicht mehr aufgerufen, siehe app/router.py).

Wird beim Programmstart EINMAL geladen und geprueft (ID eindeutig, Text nicht leer,
Pruefstatus gesetzt, jeder Platzhaltername bekannt) -- ein Fehler dort fuehrt zu einer
klaren Meldung beim Start, nicht erst bei der ersten Frage (Modulebene = Importzeitpunkt).

Diese Datei LIEST nur aus `regelwerk/`; sie veraendert dort nichts (Vorgabe: `regelwerk/`
nicht aendern, ausser durch den Menschen in der Excel selbst)."""
from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl

from app.baustein import PLATZHALTER_BESCHREIBUNG

WURZEL = Path(__file__).resolve().parent.parent
REGELKATALOG_PFAD = WURZEL / "regelwerk" / "Regelkatalog_24_KSVG_Entwurf.xlsx"

_PLATZHALTER_IM_TEXT_MUSTER = re.compile(r"\{([A-Z0-9_]+)\}")


@dataclass(frozen=True)
class Baustein:
    id: str
    reihenfolge: int
    benoetigt: tuple[str, ...]
    nicht_wenn: tuple[str, ...]
    ergebnisklasse: str | None
    text: str
    gueltig_von: int | None
    gueltig_bis: int | None

    @property
    def bedingungslos(self) -> bool:
        """Ein Baustein ganz ohne jede Bedingung (kein 'Benötigt', kein 'Nicht, wenn',
        keine Ergebnisklasse, keine Jahresgrenze) ist der Leerfall-Baustein (im
        Regelkatalog TB-16) -- er wird NUR verwendet, wenn kein anderer Baustein passt
        (siehe waehle_bausteine). Strukturell erkannt (an der fehlenden Bedingung), nicht
        ueber seine ID, damit hier keine Baustein-ID im Programmcode festgeschrieben ist."""
        return (
            not self.benoetigt and not self.nicht_wenn and self.ergebnisklasse is None
            and self.gueltig_von is None and self.gueltig_bis is None
        )


def _platzhalter_bekannt(name: str) -> bool:
    if name in PLATZHALTER_BESCHREIBUNG:
        return True
    treffer = re.match(r"^([A-Z_]+)_\d+$", name)
    return bool(treffer) and f"{treffer.group(1)}_n" in PLATZHALTER_BESCHREIBUNG


def _liste(zelle: Any) -> tuple[str, ...]:
    if not zelle:
        return ()
    return tuple(t.strip() for t in str(zelle).split(";") if t.strip())


def _lade(pfad: Path = REGELKATALOG_PFAD) -> list[Baustein]:
    wb = openpyxl.load_workbook(pfad, data_only=True, read_only=True)
    if "Textbausteine" not in wb.sheetnames:
        raise RuntimeError(f"Regelkatalog ({pfad.name}): Blatt 'Textbausteine' fehlt.")
    ws = wb["Textbausteine"]
    bausteine: list[Baustein] = []
    gesehene_ids: set[str] = set()
    for zeile in ws.iter_rows(min_row=2, values_only=True):
        if not zeile or not zeile[0]:
            continue
        (baustein_id, reihenfolge, benoetigt_roh, nicht_wenn_roh, ergebnisklasse,
         text, von, bis, pruefstatus) = zeile[:9]
        baustein_id = str(baustein_id)
        if baustein_id in gesehene_ids:
            raise RuntimeError(f"Regelkatalog, Blatt 'Textbausteine': doppelte ID '{baustein_id}'.")
        gesehene_ids.add(baustein_id)
        if not text or not str(text).strip():
            raise RuntimeError(f"Regelkatalog, Blatt 'Textbausteine', {baustein_id}: Text ist leer.")
        if not pruefstatus or not str(pruefstatus).strip():
            raise RuntimeError(f"Regelkatalog, Blatt 'Textbausteine', {baustein_id}: Prüfstatus fehlt.")
        benoetigt = _liste(benoetigt_roh)
        nicht_wenn = _liste(nicht_wenn_roh)
        for name in (*benoetigt, *nicht_wenn):
            if not _platzhalter_bekannt(name):
                raise RuntimeError(
                    f"Regelkatalog, Blatt 'Textbausteine', {baustein_id}: unbekannter Platzhalter '{name}'."
                )
        for platzhalter in _PLATZHALTER_IM_TEXT_MUSTER.findall(str(text)):
            if not _platzhalter_bekannt(platzhalter):
                raise RuntimeError(
                    f"Regelkatalog, Blatt 'Textbausteine', {baustein_id}: "
                    f"unbekannter Platzhalter '{{{platzhalter}}}' im Text."
                )
        bausteine.append(Baustein(
            id=baustein_id,
            reihenfolge=int(reihenfolge) if reihenfolge is not None else 0,
            benoetigt=benoetigt, nicht_wenn=nicht_wenn,
            ergebnisklasse=str(ergebnisklasse) if ergebnisklasse else None,
            text=str(text),
            gueltig_von=int(von) if von else None,
            gueltig_bis=int(bis) if bis else None,
        ))
    bausteine.sort(key=lambda b: (b.reihenfolge, b.id))
    return bausteine


ALLE_BAUSTEINE: list[Baustein] = _lade()


def _passt(baustein: Baustein, erlaubte_namen: set[str], ergebnisklasse: str, jahr: int | None) -> bool:
    if not all(n in erlaubte_namen for n in baustein.benoetigt):
        return False
    if any(n in erlaubte_namen for n in baustein.nicht_wenn):
        return False
    if baustein.ergebnisklasse is not None and baustein.ergebnisklasse != ergebnisklasse:
        return False
    if baustein.gueltig_von is not None and (jahr is None or jahr < baustein.gueltig_von):
        return False
    if baustein.gueltig_bis is not None and (jahr is None or jahr > baustein.gueltig_bis):
        return False
    return True


@dataclass(frozen=True)
class Auswahl:
    texte: list[str]
    leerfall: bool  # True, wenn KEIN bedingter Baustein passte und nur der Leerfall griff


def waehle_bausteine(
    werte: dict[str, str], ergebnisklasse: str, jahr: int | None, bausteine: list[Baustein] | None = None
) -> Auswahl:
    """Liefert die Rohtexte (mit `{PLATZHALTERN}`) aller passenden Bausteine, sortiert nach
    Reihenfolge/ID. Passt kein bedingter Baustein, liefert diese Funktion den
    bedingungslosen Leerfall-Baustein (Regelkatalog TB-16) als einzigen Eintrag und
    markiert `leerfall=True` -- app/router.py protokolliert das als "Baustein fehlt"
    (anonymisiert), gibt aber trotzdem Ergebnisklasse und Schlusssatz aus."""
    quelle = ALLE_BAUSTEINE if bausteine is None else bausteine
    erlaubte_namen = set(werte.keys())
    bedingte = [b for b in quelle if not b.bedingungslos and _passt(b, erlaubte_namen, ergebnisklasse, jahr)]
    if bedingte:
        return Auswahl(texte=[b.text for b in bedingte], leerfall=False)
    leerfall = next((b for b in quelle if b.bedingungslos), None)
    return Auswahl(texte=[leerfall.text] if leerfall else [], leerfall=True)
