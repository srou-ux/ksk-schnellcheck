"""Quellenindex bauen (RAG, Auftrag Teil 2, Punkt 1 und 'Index-Aktualisierung').

Befehl (Windows wie Linux):  python -m app.rag.index
    --neu                     alles neu bauen (sonst nur geänderte Quellen)
    --pruefung-erledigt R-12  Regel nach menschlicher Prüfung wieder freigeben

Liest:
- regelwerk/Regelkatalog_24_KSVG_Entwurf.xlsx (Blätter Regeln, Regelregister, Schwellen nach Jahr,
  Abgabesätze, Tätigkeitskatalog) -> Stücke der Art 'Regel'/'Register',
- quellen/index.csv + app/rag/quellen_metadaten.csv (Gültigkeit, Stufe, Kurzname, Zerlegung)
  -> Stücke aus den Quelltexten (PDF/TXT).
Schreibt: SQLite mit FTS5 (app/_lokal/rag.sqlite, nicht im Repo) und die Änderungsliste
(aufgaben/RAG_Aenderungen.md). Fassung (A/B) und Gültigkeit kommen aus den Daten: die Jahre
einer Quelle aus quellen_metadaten.csv, die Fassung je Jahr aus dem Register (kern/parameter.json).
Nichts unter quellen/ oder regelwerk/ wird verändert."""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import json
import re
import sqlite3
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from app.rag.zerlegen import KONFIG, Stueck, zerlege_datei

WURZEL = Path(__file__).resolve().parents[2]


@dataclass
class Basis:
    """Wo die Daten liegen -- für Tests mit eigener Test-Quellenbasis austauschbar."""
    wurzel: Path = WURZEL
    quellen: Path | None = None
    regelkatalog: Path | None = None
    metadaten: Path | None = None
    index: Path | None = None
    aenderungen: Path | None = None
    jahre_fassung: dict[int, str] | None = None  # None = aus dem Register (kern)

    def __post_init__(self) -> None:
        self.quellen = self.quellen or self.wurzel / "quellen"
        self.regelkatalog = self.regelkatalog or self.wurzel / "regelwerk" / "Regelkatalog_24_KSVG_Entwurf.xlsx"
        self.metadaten = self.metadaten or Path(__file__).resolve().parent / "quellen_metadaten.csv"
        self.index = self.index or self.wurzel / KONFIG["index_pfad"]
        self.aenderungen = self.aenderungen or self.wurzel / KONFIG["aenderungen_pfad"]

    def fassungen(self) -> dict[int, str]:
        if self.jahre_fassung is not None:
            return dict(self.jahre_fassung)
        from kern import rechnen as K

        return {int(j): str(v.get("fassung")) for j, v in K.P.get("schwelle", {}).items() if v and v.get("fassung")}


SCHEMA = """
CREATE TABLE IF NOT EXISTS quellen (datei TEXT PRIMARY KEY, sha256 TEXT, kurzname TEXT, art TEXT,
    stufe INTEGER, gilt_von INTEGER, gilt_bis INTEGER, stuecke INTEGER, gebaut TEXT, inhalt TEXT);
CREATE TABLE IF NOT EXISTS stuecke (id INTEGER PRIMARY KEY, datei TEXT, quelle_id TEXT, art TEXT,
    stufe INTEGER, fundstelle TEXT, titel TEXT, text TEXT, gilt_von INTEGER, gilt_bis INTEGER,
    fassung TEXT, seite INTEGER, regel_ids TEXT, status TEXT, nachrangig INTEGER DEFAULT 0);
CREATE INDEX IF NOT EXISTS stuecke_datei ON stuecke(datei);
CREATE VIRTUAL TABLE IF NOT EXISTS fts USING fts5(fundstelle, titel, text, regel_ids,
    content='stuecke', content_rowid='id', tokenize="unicode61 remove_diacritics 2");
CREATE TABLE IF NOT EXISTS regel_status (regel_id TEXT PRIMARY KEY, status TEXT, grund TEXT, seit TEXT);
CREATE TABLE IF NOT EXISTS vektoren (id INTEGER PRIMARY KEY, vektor TEXT);
CREATE TABLE IF NOT EXISTS info (schluessel TEXT PRIMARY KEY, wert TEXT);
"""

REGELKATALOG_SCHLUESSEL = "regelwerk:Regelkatalog"
# Stand der Zerlege-/Index-Logik. Ändert sich der Code, wird der Index einmal ganz neu gebaut.
INDEX_VERSION = "7"  # 7 (16.10.2026): Fassung "alle"/"Jahresabhängig" und "nicht nur gelegentlich" (app/quellen.py)


def _sha(pfad: Path) -> str:
    return hashlib.sha256(pfad.read_bytes()).hexdigest()


def _int(x: Any) -> int | None:
    try:
        return int(str(x).strip()) if str(x).strip() not in ("", "None") else None
    except ValueError:
        return None


def _fassung_text(von: int | None, bis: int | None, fassungen: dict[int, str]) -> str:
    """Fassungen (A/B …) der Register-Jahre im Gültigkeitsbereich -- aus den Daten, kein Festwert."""
    werte = sorted({f for j, f in fassungen.items() if (von is None or j >= von) and (bis is None or j <= bis)})
    return ";".join(werte)


def lies_metadaten(basis: Basis) -> dict[str, dict[str, str]]:
    """Liest app/rag/quellen_metadaten.csv. Verträgt das Speichern in Excel unter Windows
    (UTF-8 mit BOM, Semikolon als Trennzeichen)."""
    text = Path(basis.metadaten).read_text(encoding="utf-8-sig")
    kopf = text.splitlines()[0] if text else ""
    trenner = ";" if kopf.count(";") > kopf.count(",") else ","
    zeilen = list(csv.DictReader(text.splitlines(), delimiter=trenner))
    if zeilen and "datei" not in zeilen[0]:
        raise RuntimeError(f"{basis.metadaten}: Spalte 'datei' fehlt (Kopfzeile: {kopf[:80]!r}).")
    return {z["datei"]: z for z in zeilen if z.get("datei")}


def lies_index_csv(basis: Basis) -> dict[str, dict[str, str]]:
    """Nur Dateiname (Spalte 1) und Titel werden gebraucht; manche Zeilen haben verschobene
    Spalten (siehe aufgaben/FUNDE_Chatbot.md), deshalb über die Position gelesen."""
    with open(basis.quellen / "index.csv", encoding="utf-8", newline="") as f:
        zeilen = list(csv.reader(f))
    return {z[0]: {"titel": z[1] if len(z) > 1 else ""} for z in zeilen[1:] if z and z[0]}


# --------------------------------------------------------------------------- Regelkatalog

def _blatt(wb, name: str):
    for n in wb.sheetnames:
        if n.strip().lower() == name.lower():
            return wb[n]
    return None


def _jahre_aus_fassung(gilt: str | None, fassungen: dict[int, str]) -> tuple[int | None, int | None, str, bool]:
    from app.quellen import _fassung_passt

    jahre = sorted(j for j, f in fassungen.items() if _fassung_passt(gilt, f))
    if not jahre:
        return None, None, "", False
    alle = sorted(fassungen)
    von = None if jahre[0] == alle[0] else jahre[0]
    bis = None if jahre[-1] == alle[-1] else jahre[-1]
    return von, bis, ";".join(sorted({fassungen[j] for j in jahre})), True


def regel_stuecke(basis: Basis) -> tuple[list[Stueck], dict[str, list[str]]]:
    """Stücke aus dem Regelkatalog und die Zuordnung Quelldatei -> Regel-IDs."""
    import openpyxl

    wb = openpyxl.load_workbook(basis.regelkatalog, read_only=True, data_only=True)
    fassungen = basis.fassungen()
    kennz = [k.lower() for k in KONFIG["regel_stufe2_kennzeichen"]]
    datei = REGELKATALOG_SCHLUESSEL
    erg: list[Stueck] = []
    quelle_zu_regeln: dict[str, list[str]] = {}
    ws = _blatt(wb, "Regeln")
    for z in ws.iter_rows(min_row=2, values_only=True):
        if not z or not z[0]:
            continue
        rid, gruppe, norm, vor, schwelle, gilt, folge, _ein, fund, qdat = (list(z) + [None] * 12)[:10]
        vorbehalt = (list(z) + [None] * 12)[10]
        von, bis, fassung, passt = _jahre_aus_fassung(str(gilt or ""), fassungen)
        text = " ".join(str(t) for t in (gruppe, norm, vor, schwelle, folge, vorbehalt) if t)
        klein = f"{fund or ''} {vorbehalt or ''}".lower()
        stufe = 2 if any(k in klein for k in kennz) and "§" not in str(fund or "").split(";")[0] else 1
        erg.append(Stueck(datei=datei, quelle_id="Regelkatalog", art="Regel", stufe=stufe,
                          fundstelle=str(fund or "").strip(), text=text, titel=f"{rid} {gruppe or ''}".strip(),
                          gilt_von=von, gilt_bis=bis, fassung=fassung, regel_ids=[str(rid)],
                          status="aktiv" if passt else "fassung_unklar"))
        for q in re.split(r"[;,\n]", str(qdat or "")):
            q = q.strip().removeprefix("quellen/")
            if q:
                quelle_zu_regeln.setdefault(q, []).append(str(rid))
    ws = _blatt(wb, "Regelregister")
    if ws is not None:
        for z in ws.iter_rows(min_row=2, values_only=True):
            if not z or not z[0]:
                continue
            schl, wert, einheit, op, von, bis, bed, rids, fund = (list(z) + [None] * 10)[:9]
            ids = [r.strip() for r in str(rids or "").split(";") if r.strip()]
            v, b = _int(von), _int(bis)
            erg.append(Stueck(datei=datei, quelle_id="Regelregister", art="Register", stufe=1,
                              fundstelle=str(fund or "").strip(), titel=str(schl),
                              text=f"{bed or ''} ({schl}: {op or ''} {wert} {einheit or ''})",
                              gilt_von=v, gilt_bis=b, fassung=_fassung_text(v, b, fassungen), regel_ids=ids))
    for blatt, art_text in (("Schwellen nach Jahr", "Schwelle"), ("Abgabesätze", "Abgabesatz")):
        ws = _blatt(wb, blatt)
        if ws is None:
            continue
        for z in ws.iter_rows(min_row=2, values_only=True):
            jahr = _int(z[0]) if z else None
            if jahr is None:
                continue
            fund = z[3] if art_text == "Schwelle" else z[2]
            wert = z[2] if art_text == "Schwelle" else z[1]
            erg.append(Stueck(datei=datei, quelle_id=blatt, art="Register", stufe=1, fundstelle=str(fund or "").strip(),
                              titel=f"{art_text} {jahr}", text=f"{art_text} für das Kalenderjahr {jahr}: {wert}",
                              gilt_von=jahr, gilt_bis=jahr, fassung=fassungen.get(jahr, "")))
    ws = _blatt(wb, "Tätigkeitskatalog")
    if ws is not None:
        for z in ws.iter_rows(min_row=2, values_only=True):
            if not z or not z[0]:
                continue
            erg.append(Stueck(datei=datei, quelle_id="Tätigkeitskatalog", art="Register", stufe=2,
                              fundstelle=str(z[2] or "").strip(), titel=str(z[0]),
                              text=f"Tätigkeit laut Katalog: {z[0]}. {z[1] or ''}".strip(),
                              fassung=_fassung_text(None, None, fassungen)))
    return erg, quelle_zu_regeln


# --------------------------------------------------------------------------- Quelltexte

def quell_stuecke(datei: str, meta: dict[str, str], titel: str, basis: Basis) -> list[Stueck]:
    fassungen = basis.fassungen()
    von, bis = _int(meta.get("gilt_von")), _int(meta.get("gilt_bis"))
    teile = zerlege_datei(basis.quellen / datei, meta["kurzname"], meta.get("zerlegung") or "abschnitt")
    return [Stueck(datei=datei, quelle_id=meta["kurzname"], art=meta["art"], stufe=int(meta["stufe"]),
                   fundstelle=fund, text=text, titel=titel, gilt_von=von, gilt_bis=bis,
                   fassung=_fassung_text(von, bis, fassungen), seite=seite,
                   nachrangig=(meta.get("nachrangig") or "").strip().lower() == "ja")
            for fund, text, seite in teile if len(text) >= 3]


def _verknuepfe_regeln(stuecke: list[Stueck], regeln: list[Stueck]) -> None:
    """Ein Quellenstück bekommt die Regel-IDs, deren Fundstelle genau diese Stelle nennt
    (z. B. 'KSVG § 24 Abs. 2' in 'KSVG § 24 Abs. 2 Satz 1 Nr. 1')."""
    def norm(t: str) -> str:
        return re.sub(r"\s+", " ", t.replace("Satz", "S.").lower()).strip()

    nach_fund = [(norm(r.fundstelle), r.regel_ids[0], r.gilt_von, r.gilt_bis)
                 for r in regeln if r.art == "Regel" and r.fundstelle and r.status == "aktiv"]

    def ueberlappt(v1, b1, v2, b2) -> bool:
        return (v1 is None or b2 is None or v1 <= b2) and (b1 is None or v2 is None or b1 >= v2)

    for s in stuecke:
        kern = norm(re.split(r", S\. \d+$", s.fundstelle)[0])
        if not kern or kern == norm(s.quelle_id):
            s.regel_ids = []
            continue
        muster = re.compile(re.escape(kern) + r"(?![\da-z])")
        # nur Regeln, deren Gültigkeit sich mit der Quelle überschneidet (keine Fassung-A-Regel an
        # einem Text, der erst ab Fassung B gilt)
        s.regel_ids = sorted({rid for f, rid, v, b in nach_fund
                              if muster.search(f) and ueberlappt(v, b, s.gilt_von, s.gilt_bis)})


# --------------------------------------------------------------------------- Bauen

def _verbinde(pfad: Path) -> sqlite3.Connection:
    pfad.parent.mkdir(parents=True, exist_ok=True)
    db = sqlite3.connect(str(pfad))
    db.executescript(SCHEMA)
    return db


def _schreibe_stuecke(db: sqlite3.Connection, stuecke: list[Stueck]) -> None:
    for s in stuecke:
        if not s.fassung and s.status == "aktiv":
            s.status = "fassung_unklar"  # kein Register-Jahr im Gültigkeitsbereich: nie geliefert
        cur = db.execute(
            "INSERT INTO stuecke (datei, quelle_id, art, stufe, fundstelle, titel, text, gilt_von, gilt_bis,"
            " fassung, seite, regel_ids, status, nachrangig) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (s.datei, s.quelle_id, s.art, s.stufe, s.fundstelle, s.titel, s.text, s.gilt_von, s.gilt_bis,
             s.fassung, s.seite, ";".join(s.regel_ids), s.status, int(s.nachrangig)))
        db.execute("INSERT INTO fts (rowid, fundstelle, titel, text, regel_ids) VALUES (?,?,?,?,?)",
                   (cur.lastrowid, s.fundstelle, s.titel, s.text, " ".join(s.regel_ids)))


def _loesche_datei(db: sqlite3.Connection, datei: str) -> None:
    for (rid, fund, titel, text, regeln) in db.execute(
            "SELECT id, fundstelle, titel, text, regel_ids FROM stuecke WHERE datei=?", (datei,)).fetchall():
        db.execute("INSERT INTO fts (fts, rowid, fundstelle, titel, text, regel_ids) VALUES ('delete',?,?,?,?,?)",
                   (rid, fund, titel, text, (regeln or "").replace(";", " ")))
        db.execute("DELETE FROM vektoren WHERE id=?", (rid,))
    db.execute("DELETE FROM stuecke WHERE datei=?", (datei,))
    db.execute("DELETE FROM quellen WHERE datei=?", (datei,))


# Nur diese Einstellungen bestimmen, WIE zerlegt wird. Such-Einstellungen (Gewichte, Schwellen)
# dürfen keinen Neuaufbau auslösen -- sonst müssten alle Embeddings neu bezahlt werden.
_ZERLEGE_SCHLUESSEL = ("rauschen", "paragraf_kopf", "absatz_kopf", "randnummer_zeile", "randnummer_klammer",
                       "urteil_leitsatz", "abschnitt_kopf", "abschnitt_kopf_max_zeichen", "abkuerzungen_ohne_satzende",
                       "stueck_max_zeichen", "abschnitt_bezeichnung", "paragraf_zeichen", "regel_stufe2_kennzeichen")


def _konfig_sha() -> str:
    teil = {k: KONFIG.get(k) for k in _ZERLEGE_SCHLUESSEL}
    return hashlib.sha256(json.dumps(teil, sort_keys=True, ensure_ascii=False).encode("utf-8")
                          + INDEX_VERSION.encode()).hexdigest()


def _quelle_sha(pfad: Path, meta: dict[str, str]) -> str:
    """Prüfsumme aus Datei, ihrer Metadatenzeile und der Zerlege-Konfiguration: ändert sich eins
    davon (z. B. gilt_von), wird die Quelle neu zerlegt."""
    zeile = json.dumps({k: meta.get(k) for k in sorted(meta) if k}, sort_keys=True, ensure_ascii=False)
    return hashlib.sha256(pfad.read_bytes() + zeile.encode("utf-8") + _konfig_sha().encode()).hexdigest()


def _alter_stand(pfad: Path) -> tuple[dict[str, tuple[str, str]], list[tuple], str | None]:
    """({Quelle: (Prüfsumme gesamt, Prüfsumme Inhalt)}, regel_status-Zeilen, Version) eines
    vorhandenen Index. Ältere Index-Stände ohne Spalte 'inhalt' hatten als Prüfsumme nur den
    Dateiinhalt -- dann gilt diese als Inhalt."""
    if not pfad.exists():
        return {}, [], None
    db = sqlite3.connect(str(pfad))
    try:
        db.executescript(SCHEMA)
        spalten = {r[1] for r in db.execute("PRAGMA table_info(quellen)")}
        sql = "SELECT datei, sha256, inhalt FROM quellen" if "inhalt" in spalten else "SELECT datei, sha256, sha256 FROM quellen"
        alt = {d: (s, i or s) for d, s, i in db.execute(sql)}
        status = db.execute("SELECT regel_id, status, grund, seit FROM regel_status").fetchall()
        v = db.execute("SELECT wert FROM info WHERE schluessel='version'").fetchone()
        return alt, status, (v[0] if v else None)
    finally:
        db.close()


def _ersetze(neu_pfad: Path, ziel: Path) -> None:
    """Fertig gebauten Index in einem Schritt einsetzen (os.replace). Unter Windows kann die Datei
    kurz gesperrt sein, solange ein Abruf sie offen hat -- dann einige Male erneut versuchen."""
    import os
    import time

    for versuch in range(20):
        try:
            os.replace(neu_pfad, ziel)
            return
        except PermissionError:
            time.sleep(0.25 * (versuch + 1))
    os.replace(neu_pfad, ziel)


def baue(basis: Basis | None = None, neu: bool = False, ausgabe=print, liste_schreiben: bool = False) -> dict[str, Any]:
    """Baut den Index (nur Geändertes, außer `neu`). Gibt die Änderungsliste zurück. Die Datei
    aufgaben/RAG_Aenderungen.md schreibt nur der Befehl (liste_schreiben=True) -- automatische
    Aufbauten (App-Start, Tests) verändern nichts im Repo.
    Ganz neu (Befehl --neu oder geänderte Index-Version) wird in eine Nebendatei gebaut und erst
    fertig eingesetzt; laufende Abrufe lesen bis dahin den alten Stand. Die Liste der Regeln auf
    'prüfen' wird übernommen, Änderungen an Quellen werden auch dabei erkannt."""
    basis = basis or Basis()
    alt, status_alt, version = _alter_stand(basis.index)
    if version != INDEX_VERSION and alt:
        neu = True
    if not neu:
        return _baue_in(basis, basis.index, alt, [], ausgabe, liste_schreiben)
    nebendatei = basis.index.with_name(basis.index.stem + ".neu" + basis.index.suffix)
    if nebendatei.exists():
        nebendatei.unlink()
    aenderung = _baue_in(basis, nebendatei, {}, status_alt, ausgabe, liste_schreiben, vergleich=alt)
    _ersetze(nebendatei, basis.index)
    return aenderung


def _baue_in(basis: Basis, pfad: Path, alt: dict[str, tuple[str, str]], status_alt: list[tuple], ausgabe,
             liste_schreiben: bool, vergleich: dict[str, tuple[str, str]] | None = None) -> dict[str, Any]:
    db = _verbinde(pfad)
    try:
        aenderung = _baue_db(db, basis, alt, status_alt, vergleich if vergleich is not None else alt)
        db.commit()
        import os

        if os.environ.get("RAG_EMBEDDINGS") == "1":
            from app.rag.embeddings import ergaenze_vektoren

            aenderung["vektoren_neu"] = ergaenze_vektoren(db)
            db.commit()
    finally:
        db.close()
    if liste_schreiben:
        _schreibe_aenderungsliste(basis, aenderung, aenderung["_zeit"])
    ausgabe(f"RAG-Index: {basis.index} -- neu {len(aenderung['neu'])}, geändert {len(aenderung['geaendert'])}, "
            f"entfernt {len(aenderung['entfernt'])}, unverändert {aenderung['unveraendert']}, "
            f"Regeln auf 'prüfen': {len(aenderung['regeln_pruefen'])}, ohne Metadaten: {len(aenderung['ohne_metadaten'])}")
    return aenderung


def _baue_db(db: sqlite3.Connection, basis: Basis, alt: dict[str, tuple[str, str]], status_alt: list[tuple],
             vergleich: dict[str, tuple[str, str]]) -> dict[str, Any]:
    """`alt`: was in DIESER Datenbank schon steht (inkrementell). `vergleich`: Stand vor dem Aufbau
    (für Änderungsliste und 'prüfen'; bei einem Neuaufbau der alte Index)."""
    jetzt = dt.datetime.now().isoformat(timespec="seconds")
    if status_alt:
        db.executemany("INSERT OR REPLACE INTO regel_status VALUES (?,?,?,?)", status_alt)
    erstbau = not vergleich
    meta = lies_metadaten(basis)
    index_csv = lies_index_csv(basis)
    aenderung: dict[str, Any] = {"neu": [], "geaendert": [], "entfernt": [], "unveraendert": 0, "_zeit": jetzt,
                                 "neu_zerlegt": [], "ohne_metadaten": [], "nicht_aufgenommen": [],
                                 "regeln_pruefen": [], "duplikate": []}

    def melde(datei: str, sha: str, inhalt: str) -> None:
        """'geändert' = Inhalt der Datei anders (setzt Regeln auf prüfen); nur Metadaten/Zerlegung
        anders = 'neu zerlegt' (kein prüfen, der Quelltext ist derselbe)."""
        frueher = vergleich.get(datei)
        if frueher is None:
            aenderung["neu"].append(datei)
        elif frueher[1] != inhalt:
            aenderung["geaendert"].append(datei)
        elif frueher[0] != sha:
            aenderung["neu_zerlegt"].append(datei)
        else:
            aenderung["unveraendert"] += 1

    # Regelkatalog (+ Register-Jahre) -- immer zuerst, die Quellen verknüpfen sich damit
    kern_param = basis.wurzel / "kern" / "parameter.json"
    sha_regel = hashlib.sha256(basis.regelkatalog.read_bytes()
                               + (kern_param.read_bytes() if kern_param.exists() else b"")
                               + json.dumps(basis.jahre_fassung, sort_keys=True).encode()
                               + _konfig_sha().encode()).hexdigest()
    regeln, quelle_zu_regeln = regel_stuecke(basis)
    inhalt_regel = hashlib.sha256(basis.regelkatalog.read_bytes()).hexdigest()
    regelkatalog_neu = (alt.get(REGELKATALOG_SCHLUESSEL) or ("", ""))[0] != sha_regel
    if regelkatalog_neu:
        _loesche_datei(db, REGELKATALOG_SCHLUESSEL)
        _schreibe_stuecke(db, regeln)
        db.execute("INSERT OR REPLACE INTO quellen VALUES (?,?,?,?,?,?,?,?,?,?)",
                   (REGELKATALOG_SCHLUESSEL, sha_regel, "Regelkatalog", "Regel", 1, None, None, len(regeln), jetzt, inhalt_regel))
    melde(REGELKATALOG_SCHLUESSEL, sha_regel, inhalt_regel)

    gesehen_inhalt: dict[str, str] = {}
    vorhanden: set[str] = {REGELKATALOG_SCHLUESSEL}
    for datei in index_csv:
        pfad = basis.quellen / datei
        m = meta.get(datei)
        if m is None:
            aenderung["ohne_metadaten"].append(datei)
            continue
        if (m.get("aufnehmen") or "ja").strip().lower() != "ja":
            aenderung["nicht_aufgenommen"].append(f"{datei} ({m.get('bemerkung') or 'aufnehmen = nein'})")
            continue
        if not pfad.exists():
            continue  # fehlt die Datei, gilt sie unten als entfernt
        inhalt = _sha(pfad)
        if inhalt in gesehen_inhalt:
            aenderung["duplikate"].append(f"{datei} = {gesehen_inhalt[inhalt]}")
            continue
        gesehen_inhalt[inhalt] = datei
        vorhanden.add(datei)
        sha = _quelle_sha(pfad, m)
        melde(datei, sha, inhalt)
        if (alt.get(datei) or ("", ""))[0] == sha:
            continue
        if datei in alt:
            _loesche_datei(db, datei)
        stuecke = quell_stuecke(datei, m, index_csv[datei]["titel"], basis)
        _verknuepfe_regeln(stuecke, regeln)
        _schreibe_stuecke(db, stuecke)
        db.execute("INSERT OR REPLACE INTO quellen VALUES (?,?,?,?,?,?,?,?,?,?)",
                   (datei, sha, m["kurzname"], m["art"], int(m["stufe"]), _int(m.get("gilt_von")),
                    _int(m.get("gilt_bis")), len(stuecke), jetzt, inhalt))
    for datei in set(alt) - vorhanden:
        _loesche_datei(db, datei)
    aenderung["entfernt"] = sorted(set(vergleich) - vorhanden)

    # Regelkatalog geändert: auch die unveränderten Quellen neu mit den Regeln verknüpfen
    if regelkatalog_neu:
        _verknuepfe_bestand(db, regeln)

    # Geänderte oder entfernte Quelle -> verknüpfte Regeln 'prüfen' (bleibt bis zur Freigabe)
    if not erstbau:
        for datei in aenderung["geaendert"] + aenderung["entfernt"]:
            if datei == REGELKATALOG_SCHLUESSEL:
                continue
            art = "entfernt" if datei in aenderung["entfernt"] else "geändert"
            for rid in sorted(set(quelle_zu_regeln.get(datei, []))):
                db.execute("INSERT OR REPLACE INTO regel_status VALUES (?,?,?,?)",
                           (rid, "prüfen", f"Quelle {art}: {datei}", jetzt))
                aenderung["regeln_pruefen"].append(f"{rid} (Quelle {art}: {datei})")
    db.execute("INSERT OR REPLACE INTO info VALUES ('gebaut', ?)", (jetzt,))
    db.execute("INSERT OR REPLACE INTO info VALUES ('version', ?)", (INDEX_VERSION,))
    return aenderung


def _verknuepfe_bestand(db: sqlite3.Connection, regeln: list[Stueck]) -> None:
    zeilen = db.execute("SELECT id, datei, quelle_id, fundstelle, titel, text, gilt_von, gilt_bis, regel_ids FROM stuecke "
                        "WHERE art NOT IN ('Regel','Register')").fetchall()
    stuecke = [Stueck(datei=z[1], quelle_id=z[2], art="", stufe=1, fundstelle=z[3], text=z[5], titel=z[4],
                      gilt_von=z[6], gilt_bis=z[7]) for z in zeilen]
    _verknuepfe_regeln(stuecke, regeln)
    for z, s in zip(zeilen, stuecke):
        neu_ids = ";".join(s.regel_ids)
        if neu_ids == (z[8] or ""):
            continue
        db.execute("INSERT INTO fts (fts, rowid, fundstelle, titel, text, regel_ids) VALUES ('delete',?,?,?,?,?)",
                   (z[0], z[3], z[4], z[5], (z[8] or "").replace(";", " ")))
        db.execute("INSERT INTO fts (rowid, fundstelle, titel, text, regel_ids) VALUES (?,?,?,?,?)",
                   (z[0], z[3], z[4], z[5], neu_ids.replace(";", " ")))
        db.execute("UPDATE stuecke SET regel_ids=? WHERE id=?", (neu_ids, z[0]))


def _schreibe_aenderungsliste(basis: Basis, a: dict[str, Any], jetzt: str) -> None:
    def liste(titel: str, eintraege: list[str]) -> list[str]:
        return [f"### {titel} ({len(eintraege)})", *(f"- {e}" for e in eintraege or ["–"]), ""]

    zeilen = ["# RAG – Änderungsliste des letzten Index-Aufbaus", "",
              f"Erzeugt von `python -m app.rag.index` am {jetzt}. Wird bei jedem Aufbau neu geschrieben.", "",
              f"Unverändert: {a['unveraendert']}", ""]
    zeilen += liste("Neu", a["neu"]) + liste("Geändert (Inhalt der Datei)", a["geaendert"]) + liste("Entfernt", a["entfernt"])
    zeilen += liste("Neu zerlegt (nur Metadaten oder Zerlege-Einstellungen geändert, Inhalt gleich)", a.get("neu_zerlegt", []))
    zeilen += liste("Regeln auf „prüfen“ gesetzt", a["regeln_pruefen"])
    zeilen += liste("In quellen/index.csv, aber ohne Zeile in app/rag/quellen_metadaten.csv (nicht indexiert)", a["ohne_metadaten"])
    zeilen += liste("Bewusst nicht aufgenommen", a["nicht_aufgenommen"])
    zeilen += liste("Inhaltsgleich (nur einmal indexiert)", a["duplikate"])
    basis.aenderungen.parent.mkdir(parents=True, exist_ok=True)
    basis.aenderungen.write_text("\n".join(zeilen), encoding="utf-8")


def pruefung_erledigt(regel_id: str, basis: Basis | None = None) -> bool:
    basis = basis or Basis()
    db = _verbinde(basis.index)
    n = db.execute("DELETE FROM regel_status WHERE regel_id=?", (regel_id,)).rowcount
    db.commit()
    db.close()
    return n > 0


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description="RAG-Index bauen (nur geänderte Quellen).")
    p.add_argument("--neu", action="store_true", help="alles neu bauen")
    p.add_argument("--pruefung-erledigt", metavar="REGEL_ID", help="Regel nach Prüfung wieder freigeben")
    a = p.parse_args(argv)
    if a.pruefung_erledigt:
        ok = pruefung_erledigt(a.pruefung_erledigt)
        print(f"{a.pruefung_erledigt}: {'freigegeben' if ok else 'war nicht auf prüfen'}")
        return 0
    baue(neu=a.neu, liste_schreiben=True)
    return 0


if __name__ == "__main__":
    sys.exit(main())
