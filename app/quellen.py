"""Quellenabruf (RAG, nur oeffentliche Quellen, Architektur Punkt 4).

Liest ausschliesslich lokale Dateien (`regelwerk/Regelkatalog_24_KSVG_Entwurf.xlsx`,
`quellen/index.csv`) -- kein Netzwerkzugriff zur Laufzeit. Eine Regel wird nur zurueckgegeben,
wenn Jahr/Fassung passen. Wird auch von `app/pruefung.py` genutzt, um zu prüfen, ob eine
genannte Fundstelle/Regel-ID/ein Aktenzeichen ueberhaupt in der Quellenbasis existiert.

Diese Datei LIEST nur aus `regelwerk/` und `quellen/`; sie veraendert dort nichts
(Vorgabe: kern/, regelwerk/, quellen/ nicht aendern).
"""
from __future__ import annotations

import csv
import functools
import re
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

import openpyxl

WURZEL = Path(__file__).resolve().parent.parent
REGELKATALOG_PFAD = WURZEL / "regelwerk" / "Regelkatalog_24_KSVG_Entwurf.xlsx"
INDEX_CSV_PFAD = WURZEL / "quellen" / "index.csv"

# Erwartete Spaltenreihenfolge im Blatt "Regeln" (per Position, nicht per moeglicherweise
# falsch codiertem Spaltennamen -- die Xlsx enthaelt Umlaute).
_REGEL_SPALTEN = [
    "regel_id", "gruppe", "norm", "voraussetzung", "schwelle_zeitraum",
    "gilt_fuer_fassung", "rechtsfolge", "benoetigte_eingaben", "fundstelle",
    "quelldateien", "vorbehalt", "pruefstatus",
]
_SCHWELLE_SPALTEN = ["kalenderjahr", "fassung", "schwelle", "fundstelle", "beleg", "vorbehalt"]
_ABGABESATZ_SPALTEN = ["kalenderjahr", "satz_prozent", "fundstelle", "beleg", "vorbehalt"]


def fassung_fuer_jahr(jahr: int) -> str | None:
    """Fassung A (bis 31.12.2022) oder B (ab 01.01.2023, 8. SGB IV-AendG) -- kommt aus dem
    Kern-Register (`kern/parameter.json`, Blatt 'Schwellen nach Jahr'), NICHT aus einem
    Festwert wie 'jahr <= 2022' im Code (Teil 1f, Befund D; CLAUDE.md Regel 2). Deckt das
    Register das Jahr nicht ab, ist die Fassung unbekannt (`None`) -- Aufrufer behandeln das
    wie jede andere fehlende Angabe (siehe quellen/NICHT_GEFUNDEN.md und
    regelwerk/Regelkatalog_24_KSVG_Entwurf.xlsx, Blatt 'Offene Punkte')."""
    from kern import rechnen as K

    schwelle = K.schwelle(jahr)
    return schwelle["fassung"] if schwelle else None


def verfuegbare_jahre() -> list[int]:
    """Alle Kalenderjahre, die das Register (kern/parameter.json, Blatt 'Schwellen nach
    Jahr') abdeckt -- fuer Auswahlknoepfe bei der Jahresfrage (Teil 1f, Abschnitt J.1).
    Kein Festwert im Code: direkt aus dem Register gelesen, nicht aus einer Zahl hier."""
    from kern import rechnen as K

    return sorted(int(j) for j in K.P.get("schwelle", {}).keys())


def _blattname(wb: Any, *, exakt: str | None = None, praefix: str | None = None) -> str:
    """Findet ein Arbeitsblatt ueber Name statt Position -- Positionen sind nicht stabil,
    siehe aufgaben/FUNDE_Chatbot.md (ein neues Blatt 'Regelregister' hat alle folgenden
    Blatt-Indizes verschoben). `praefix` ist ASCII, um Umlaut-Encoding-Unsicherheiten zu
    vermeiden; `exakt` wird case-insensitiv verglichen."""
    for name in wb.sheetnames:
        n = name.strip().lower()
        if exakt is not None and n == exakt.lower():
            return name
        if praefix is not None and n.startswith(praefix.lower()):
            return name
    raise KeyError(f"Arbeitsblatt nicht gefunden (exakt={exakt!r}, praefix={praefix!r}); vorhanden: {wb.sheetnames}")


def _oeffne_arbeitsblatt(*, exakt: str | None = None, praefix: str | None = None):
    wb = openpyxl.load_workbook(REGELKATALOG_PFAD, read_only=True, data_only=True)
    return wb[_blattname(wb, exakt=exakt, praefix=praefix)]


def _zeilen_als_dicts(spalten: list[str], *, exakt: str | None = None, praefix: str | None = None) -> list[dict[str, Any]]:
    ws = _oeffne_arbeitsblatt(exakt=exakt, praefix=praefix)
    ergebnis = []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True)):
        if row[0] is None:
            continue
        eintrag = dict(zip(spalten, row))
        eintrag["_zeile"] = i + 2
        ergebnis.append(eintrag)
    return ergebnis


@functools.lru_cache(maxsize=1)
def alle_regeln() -> tuple[dict[str, Any], ...]:
    # exakt="Regeln" (nicht nur ein Praefix!), sonst wuerde das Blatt "Regelregister"
    # faelschlich mitgefunden (startswith("regeln") ist False dafuer, aber ein blosses
    # "regel"-Praefix waere falsch -- siehe Test test_blattname_regeln_nicht_regelregister).
    return tuple(_zeilen_als_dicts(_REGEL_SPALTEN, exakt="Regeln"))


@functools.lru_cache(maxsize=1)
def alle_schwellen() -> tuple[dict[str, Any], ...]:
    return tuple(_zeilen_als_dicts(_SCHWELLE_SPALTEN, exakt="Schwellen nach Jahr"))


@functools.lru_cache(maxsize=1)
def alle_abgabesaetze() -> tuple[dict[str, Any], ...]:
    return tuple(_zeilen_als_dicts(_ABGABESATZ_SPALTEN, praefix="Abgabes"))


@functools.lru_cache(maxsize=1)
def index_csv_zeilen() -> tuple[dict[str, str], ...]:
    with open(INDEX_CSV_PFAD, encoding="utf-8", newline="") as f:
        return tuple(csv.DictReader(f))


def _fassung_passt(gilt_fuer_fassung: str | None, fassung: str) -> bool:
    """Prueft, ob eine Regel fuer die gegebene Fassung (A/B) gilt. Arbeitet klausel-weise
    (getrennt durch ';'), weil Zellen wie 'Fassung A (bis 31.12.2022); B nicht geprueft'
    vorkommen: 'B' taucht dort zwar als Wort auf, ist aber durch 'nicht' ausdruecklich
    NICHT bestaetigt. Im Zweifel also nicht zuordnen (lieber 'nicht abgedeckt' als falsch
    zuordnen)."""
    if not gilt_fuer_fassung:
        return False
    for klausel in gilt_fuer_fassung.split(";"):
        klausel = klausel.strip()
        if not klausel:
            continue
        # RAG-Güte 16.10.2026: 'A (bis 2022: nicht nur gelegentlich) / B (ab 2023)' (R-115) galt als verneint, weil das
        # 'nicht' des Merkmals 'nicht nur gelegentlich' mitgelesen wurde -- die Regel wurde nie geliefert. Verneint ist eine
        # Klausel nur, wenn 'nicht' die Prüfung betrifft ('B nicht geprüft', 'noch nicht verglichen').
        verneint = re.search(r"\bnicht\s+(?:\w+\s+)?(?:gepr(?:ü|ue)ft|verglichen|best(?:ä|ae)tigt|belegt)",
                             klausel, re.IGNORECASE) is not None
        if not verneint and re.fullmatch(r"(?:alle|jahresabh(?:ä|ae)ngig)", klausel, re.IGNORECASE):
            return True  # ausdrücklich für jedes Jahr (R-20, R-21, R-89); welcher Wert gilt, sagt das Register je Jahr
        gefunden = {m.upper() if m.lower() != "aktuell" else "aktuell" for m in re.findall(r"\b(A|B|aktuell)\b", klausel, re.IGNORECASE)}
        if fassung in gefunden and not verneint:
            return True
    return False


def regeln_suchen(jahr: int, stichworte: list[str], max_treffer: int = 8) -> list[dict[str, Any]]:
    """Regeln, deren 'Gilt fuer Fassung' zum Jahr passt und die mindestens eines der
    Stichworte in Gruppe/Norm/Voraussetzung/Rechtsfolge enthalten. Reihenfolge wie im
    Regelkatalog (keine Relevanzsortierung ueber den Wortlaut hinaus)."""
    fassung = fassung_fuer_jahr(jahr)
    stichworte_klein = [s.lower() for s in stichworte if s]
    treffer = []
    for regel in alle_regeln():
        if not _fassung_passt(regel.get("gilt_fuer_fassung"), fassung):
            continue
        haystack = " ".join(
            str(regel.get(feld) or "") for feld in ("gruppe", "norm", "voraussetzung", "rechtsfolge")
        ).lower()
        if not stichworte_klein or any(s in haystack for s in stichworte_klein):
            treffer.append(regel)
        if len(treffer) >= max_treffer:
            break
    return treffer


def regel_existiert(regel_id: str) -> bool:
    return any(r.get("regel_id") == regel_id for r in alle_regeln())


def regel_by_id(regel_id: str) -> dict[str, Any] | None:
    """Direkter Lookup einer Regel-ID -- wird von app/router.py genutzt, weil `kern.rechnen`
    bei jedem Ergebnis bereits die exakt verwendeten Regel-IDs mitliefert (`Ergebnis.regeln`);
    das ist praeziser als eine Stichwortsuche."""
    for r in alle_regeln():
        if r.get("regel_id") == regel_id:
            return r
    return None


def regeln_suchen_ohne_jahr(stichworte: list[str], max_treffer: int = 8) -> list[dict[str, Any]]:
    """Wie `regeln_suchen`, aber ohne Fassungs-Filter -- fuer reine Rechtsprechungs-/
    Auslegungsfragen ohne Jahresangabe (Befund 6 aus aufgaben/FUNDE_Chatbot.md)."""
    stichworte_klein = [s.lower() for s in stichworte if s]
    if not stichworte_klein:
        return []
    treffer = []
    for regel in alle_regeln():
        haystack = " ".join(
            str(regel.get(feld) or "") for feld in ("gruppe", "norm", "voraussetzung", "rechtsfolge")
        ).lower()
        if any(s in haystack for s in stichworte_klein):
            treffer.append(regel)
        if len(treffer) >= max_treffer:
            break
    return treffer


def schwelle_fuer_jahr(jahr: int) -> dict[str, Any] | None:
    for zeile in alle_schwellen():
        if zeile.get("kalenderjahr") == jahr:
            return zeile
    return None


def abgabesatz_fuer_jahr(jahr: int) -> Decimal | None:
    for zeile in alle_abgabesaetze():
        if zeile.get("kalenderjahr") == jahr:
            wert = zeile.get("satz_prozent")
            try:
                return Decimal(str(wert))
            except (InvalidOperation, TypeError):
                return None
    return None


def dateiname_in_index(dateiname: str) -> bool:
    gesucht = dateiname.strip().lower()
    return any(gesucht in (z.get("Dateiname") or "").lower() for z in index_csv_zeilen())


_ARCHIV_KOPIE_MUSTER = re.compile(r"\s*\(Kopie\s+\d{4}-\d{2}-\d{2}\)", re.IGNORECASE)


def bereinige_fundstelle(text: str | None) -> str | None:
    """Teil 1f, Befund D: manche Fundstellen im Regelkatalog tragen eine Archiv-Kopie-
    Markierung mit festem Datum ('KSVG § 24 Abs. 2 Satz 2 (Kopie 2022-01-22)') -- das
    Archivdatum einer Quelle ist nicht ihr Rechtsstand und darf nicht wie ein fuer die
    Anfrage gueltiges Jahr wirken (siehe aufgaben/SYSTEMBESCHREIBUNG_fuer_KI.md, Abschnitt 4).
    Ersetzt die Markierung durch einen datumslosen Vorbehalt; die eigentliche Norm-Angabe
    bleibt unveraendert (nur `app/`, nicht `regelwerk/`, wird hier angepasst)."""
    if not text:
        return text
    if _ARCHIV_KOPIE_MUSTER.search(text):
        return _ARCHIV_KOPIE_MUSTER.sub(" (Stand der Quelle: Archivfassung, nicht der Rechtsstand)", text)
    return text


_AKTENZEICHEN_MUSTER = re.compile(r"[A-Z]+\s*\d+\s*[A-Z]{0,3}\s*\d+/\d+(?:\s*[A-Z])?", re.IGNORECASE)


def _normalisiere(text: str) -> str:
    return re.sub(r"\s+", " ", text.strip().lower())


def fundstelle_bekannt(text: str) -> bool:
    """Prueft, ob ein Fundstellen-Text (Regel-ID, Aktenzeichen, Norm-Zitat, KSK-Info-Nummer)
    irgendwo in der Quellenbasis vorkommt: im Regelkatalog (Fundstelle/Vorbehalt/Quelldateien)
    oder in quellen/index.csv (Dateiname/Titel/Hinweis)."""
    ziel = _normalisiere(text)
    if not ziel:
        return False

    if re.fullmatch(r"r-\d+", ziel):
        return regel_existiert(ziel.upper())

    for regel in alle_regeln():
        heu = " ".join(
            str(regel.get(f) or "") for f in ("fundstelle", "vorbehalt", "quelldateien", "norm")
        )
        if ziel in _normalisiere(heu):
            return True

    for zeile in index_csv_zeilen():
        heu = " ".join(zeile.get(k) or "" for k in ("Dateiname", "Titel", "Hinweis"))
        if ziel in _normalisiere(heu):
            return True

    return False


def aktenzeichen_in_text(text: str) -> list[str]:
    """Extrahiert aktenzeichen-aehnliche Muster (z. B. 'B 3 KS 3/21 R') aus einem Text."""
    return [m.group(0).strip() for m in _AKTENZEICHEN_MUSTER.finditer(text)]
