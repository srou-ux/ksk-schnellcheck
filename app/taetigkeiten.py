"""Tätigkeitsprüfung (Testfall 20 Positionen, 11.10.2026).

Vorher galt jede Zahlung an eine Einzelperson als Honorar für eine künstlerische Leistung --
auch Programmierung, SEO oder Datenschutzberatung kamen als sicher "Abgabepflichtig" heraus.
Jetzt wird jede zählende Position einer Tätigkeit aus dem Excel-Blatt
'Tätigkeitsabgrenzung' zugeordnet (regelwerk/Regelkatalog_24_KSVG_Entwurf.xlsx). Gruppen:

- künstlerisch: zählt wie bisher,
- nicht künstlerisch: zählt nicht (der Kern schließt sie über `kuenstlerisch: False` aus),
- Grenzfall: der Steuerberater entscheidet,
- nicht in der Liste ('unbekannt') oder Modell und feste Erkennung widersprechen sich
  ('widerspruch'): wie Grenzfall -- nie "Abgabepflichtig".

Tätigkeiten, Gruppen, Erkennungsbegriffe, Regel-IDs und Fundstellen stehen ausschließlich in
der Excel; hier steht nur das Verfahren. Das Sprachmodell darf nur eine T-ID aus dieser Liste
wählen (app/llm/mistral.py); die feste Erkennung über die Begriffe ist die Gegenprobe.
"""
from __future__ import annotations

import json
import re
from functools import lru_cache
from pathlib import Path
from typing import Any, Optional

import openpyxl

from app.quellen import REGELKATALOG_PFAD

BLATT = "Tätigkeitsabgrenzung"
UNBEKANNT = "unbekannt"
WIDERSPRUCH = "widerspruch"

_GRUPPEN = {"kuenstlerisch": "kuenstlerisch", "nicht kuenstlerisch": "nicht_kuenstlerisch", "grenzfall": "grenzfall"}


_UMLAUTE = str.maketrans({"ä": "ae", "ö": "oe", "ü": "ue", "ß": "ss"})


_SELBER = {"zaehlt mit": "zaehlt", "zaehlt nicht": "zaehlt_nicht", "steuerberater entscheidet": "unsicher"}


def _norm(text: str) -> str:
    """Klein, Umlaute als ae/oe/ue/ss (Modell und Demo-Modus schreiben sie verschieden)."""
    return " ".join(str(text).lower().translate(_UMLAUTE).split())


@lru_cache(maxsize=1)
def alle() -> tuple[dict[str, Any], ...]:
    wb = openpyxl.load_workbook(REGELKATALOG_PFAD, read_only=True, data_only=True)
    if BLATT not in wb.sheetnames:
        return ()
    eintraege = []
    for zeile in wb[BLATT].iter_rows(min_row=2, values_only=True):
        if not zeile or not zeile[0]:
            continue
        (tid, name, gruppe, vorrang, begriffe, bedingung, regeln, fundstelle, status, _bemerkung,
         selber_kuenstler, selber_regeln, selber_hinweis) = (list(zeile) + [None] * 13)[:13]
        gruppe_schluessel = _GRUPPEN.get(_norm(gruppe or ""))
        if gruppe_schluessel is None:
            continue  # unbekannte Gruppe: lieber nicht verwenden als falsch einordnen
        eintraege.append({
            "id": str(tid).strip(),
            "name": str(name or "").strip(),
            "gruppe": gruppe_schluessel,
            "vorrang": int(vorrang) if vorrang is not None else 2,
            "begriffe": [_norm(b) for b in str(begriffe or "").split(";") if b.strip()],
            "bedingung": str(bedingung or "").strip(),
            "regel_ids": [r.strip() for r in str(regeln or "").split(";") if r.strip()],
            "fundstellen": [f.strip() for f in str(fundstelle or "").split(";") if f.strip()],
            "status": str(status or "").strip(),
            # Mehrere Positionen desselben Künstlers (R-42, R-117, R-175): zählt diese Leistung
            # mit, wenn derselbe Künstler sie neben einer künstlerischen Leistung abrechnet?
            "selber_kuenstler": _SELBER.get(_norm(selber_kuenstler or ""), "unsicher"),
            "selber_regel_ids": [r.strip() for r in str(selber_regeln or "").split(";") if r.strip()],
            "selber_hinweis": str(selber_hinweis or "").strip(),
        })
    return tuple(eintraege)


def nach_id(tid: Optional[str]) -> Optional[dict[str, Any]]:
    for e in alle():
        if e["id"] == tid:
            return e
    return None


def gueltige_id(tid: Optional[str]) -> bool:
    return nach_id(tid) is not None


_SPRACHE = json.loads((Path(__file__).resolve().parent / "taetigkeit_erkennung.json").read_text(encoding="utf-8"))
_ENDUNGEN = "|".join(re.escape(_norm(e)) for e in sorted(_SPRACHE["endungen"], key=len, reverse=True) if e)
_WORTENDE = frozenset(_norm(w) for w in _SPRACHE["wortende_begriffe"])
_EIGENER_ZWECK = re.compile(r"\W{1,3}(?:" + "|".join(re.escape(_norm(x)) for x in _SPRACHE["eigener_zweck"]) + r")\b")


@lru_cache(maxsize=None)
def _muster(begriff: str) -> re.Pattern[str]:
    """Wortanfang (vor dem Begriff kein Buchstabe und keine Ziffer) -- oder, nur für die freigegebenen Kernwörter, das
    ENDE eines zusammengesetzten Worts ('werbejingle', 'hochzeitsfotografin'; app/taetigkeit_erkennung.json)."""
    anfang = r"(?<![0-9a-z])" + re.escape(begriff)
    if begriff not in _WORTENDE:
        return re.compile(anfang)
    endung = f"(?:{_ENDUNGEN})?" if _ENDUNGEN else ""
    im_wort = r"(?<=[a-z])" + re.escape(begriff) + endung + r"(?![a-z0-9])"
    return re.compile(f"{anfang}|{im_wort}")


_VERNEINUNG = re.compile(r"(?:\bohne|\bnicht|\bexkl\w*|\babzueglich|\bausser|\bkeine?)\W+(?:\w+\W+){0,1}$")


def _ist_eigener_zweck(t: str, ende: int) -> bool:
    """'Öffentlichkeitsarbeit für sich selbst' beschreibt den Zweck des Auftraggebers, nicht die Leistung."""
    folgt = t[ende:ende + int(_SPRACHE["abstand"]) + max(map(len, _SPRACHE["eigener_zweck"]))]
    m = _EIGENER_ZWECK.match(folgt)
    return bool(m) and m.start() <= int(_SPRACHE["abstand"])


_KONTAKT = tuple(_norm(w) for w in _SPRACHE.get("kontakt_woerter", []))
_NUMMER_DANACH = re.compile(r"[a-z]*\.?:?\s*\+?(\d[\d /()-]*\d)(?!\d)(?![.,]\d)(?!\s*(?:€|eur\b|euro))")


def _ist_kontaktangabe(t: str, begriff: str, ende: int) -> bool:
    """Siehe app/taetigkeit_erkennung.json, '_hinweis_kontakt': 'Telefon 030 1234567' ist keine Leistung."""
    if not begriff.startswith(_KONTAKT) or not _KONTAKT:
        return False
    m = _NUMMER_DANACH.match(t, ende)
    return bool(m) and sum(c.isdigit() for c in m.group(1)) >= 6


_EIGENE_SEITE = re.compile(_SPRACHE["eigene_seite_davor"]) if _SPRACHE.get("eigene_seite_davor") else None


def _beschreibt_auftraggeber(t: str, anfang: int) -> bool:
    """Siehe app/taetigkeit_erkennung.json, '_hinweis_eigene_seite'. Nur im selben Satz."""
    if _EIGENE_SEITE is None:
        return False
    davor = re.split(r"[.!?;\n]", t[max(0, anfang - 60):anfang])[-1]
    return bool(_EIGENE_SEITE.search(davor))


def _bejahte_stellen(t: str, begriff: str) -> list[tuple[int, int]]:
    """Stellen (Anfang, Ende), an denen der Begriff vorkommt, nicht verneint ist ('Grafikleistung (ohne Druckkosten)'
    nennt keine Druckkosten als Leistung; Livetest 11.10.2026) und nicht den eigenen Zweck des Auftraggebers meint."""
    stellen = []
    for treffer in _muster(begriff).finditer(t):
        if _VERNEINUNG.search(t[max(0, treffer.start() - 25):treffer.start()]):
            continue
        if (_ist_eigener_zweck(t, treffer.end()) or _ist_kontaktangabe(t, begriff, treffer.end())
                or _beschreibt_auftraggeber(t, treffer.start())):
            continue
        stellen.append((treffer.start(), treffer.end()))
    return stellen


def _bejaht(t: str, begriff: str) -> bool:
    return bool(_bejahte_stellen(t, begriff))


def _verneint(t: str, begriff: str) -> bool:
    """Der Begriff kommt vor, aber verneint ('keine Gestaltung')."""
    return any(_VERNEINUNG.search(t[max(0, m.start() - 25):m.start()]) for m in _muster(begriff).finditer(t))


def _treffer_ohne_eingeschlossene(t: str) -> list[dict[str, Any]]:
    """Zeilen mit bejahtem Begriff. Liegt JEDE Fundstelle einer Zeile ganz innerhalb einer längeren Fundstelle einer
    anderen Zeile ('telefon' in 'telefonakquise'), zählt nur die längere (Korpus-Lauf 15.10.2026: 'Telefonakquise'
    traf T-35 und T-67 -> Widerspruch)."""
    stellen = {e["id"]: [s for b in e["begriffe"] for s in _bejahte_stellen(t, b)] for e in alle()}
    treffer = [e for e in alle() if stellen[e["id"]]]

    def eingeschlossen(e) -> bool:
        # Nur von einer Zeile mit gleichem oder besserem Vorrang verdrängen lassen (Musiklehrer: T-71 mit Vorrang 0
        # darf nicht hinter T-18 'musiklehrer' verschwinden).
        andere = [s for x in treffer if x["id"] != e["id"] and x["vorrang"] <= e["vorrang"] for s in stellen[x["id"]]]
        return all(any(a <= s0 and s1 <= b and (b - a) > (s1 - s0) for a, b in andere) for s0, s1 in stellen[e["id"]])

    return [e for e in treffer if not eingeschlossen(e)]


def erkenne(text: Optional[str]) -> Optional[str]:
    """Feste Erkennung über die Begriffe der Excel. Liefert eine T-ID, `WIDERSPRUCH` (Treffer
    gleichen Vorrangs aus verschiedenen Gruppen) oder None (kein Treffer).

    Alltagstest AT-2026-036 (15.10.2026): 'Festival-Plakatierung (aufkleben, keine Gestaltung)' traf über den
    Begriff 'plakat' Grafikdesign (T-02, künstlerisch), obwohl der Text 'Gestaltung' -- einen Begriff derselben
    Zeile -- ausdrücklich verneint. Verneint der Text einen Begriff der gewählten Zeile, ist die Einordnung nicht
    eindeutig (WIDERSPRUCH, der Steuerberater entscheidet). Keine neue Fachaussage, nur keine stille Wahl."""
    if not text:
        return None
    t = _norm(text)
    treffer = _treffer_ohne_eingeschlossene(t)
    if not treffer:
        return None
    bester = min(e["vorrang"] for e in treffer)
    oben = [e for e in treffer if e["vorrang"] == bester]
    if len({e["gruppe"] for e in oben}) > 1:
        return WIDERSPRUCH
    if any(_verneint(t, b) for b in oben[0]["begriffe"]):
        return WIDERSPRUCH
    return oben[0]["id"]


def gruppen_im_text(text: Optional[str]) -> set[str]:
    """Alle Gruppen, deren Begriffe im Text vorkommen (ohne Vorrang) -- für die Vorsicht bei
    mehreren Positionen ohne eigene Bezeichnung (app/zuordnung.py)."""
    if not text:
        return set()
    t = _norm(text)
    return {e["gruppe"] for e in _treffer_ohne_eingeschlossene(t)}


def kombiniere(modell_id: Optional[str], erkannt: Optional[str], beschrieben: bool) -> Optional[str]:
    """Modellwahl und feste Erkennung zusammenführen. Widersprechen sich die Gruppen, wird
    nicht gewählt (WIDERSPRUCH). Ohne jede Zuordnung: UNBEKANNT, wenn die Leistung beschrieben
    ist, sonst None (dann fragt app/pflichtangaben.py nach der Leistung)."""
    modell = modell_id if gueltige_id(modell_id) else None
    if erkannt == WIDERSPRUCH:
        return WIDERSPRUCH
    if modell and erkannt and nach_id(modell)["gruppe"] != nach_id(erkannt)["gruppe"]:
        return WIDERSPRUCH
    if erkannt:
        return erkannt
    if modell and nach_id(modell)["gruppe"] == "nicht_kuenstlerisch":
        # Gesamtlauf b87a9bb (HT-2026-074 FALSCH): 'Kunstdrucke' -- das Modell allein wählte T-40 (Druck, nicht
        # künstlerisch), kein Begriff der Excel stand im Text; daraus wurde still 'nicht abgabepflichtig'. Ein
        # AUSSCHLUSS braucht einen Beleg im Text (Begriff aus dem Blatt 'Tätigkeitsabgrenzung'); sonst ist die
        # Leistung nicht sicher einzuordnen (Steuerberater entscheidet, beide Beträge werden gezeigt).
        return UNBEKANNT
    if modell:
        return modell
    if beschrieben or modell_id == UNBEKANNT:
        return UNBEKANNT
    return None


def gruppe(tid: Optional[str]) -> Optional[str]:
    """'kuenstlerisch', 'nicht_kuenstlerisch', 'unsicher' (Grenzfall, unbekannt, Widerspruch)
    oder None (noch keine Angabe)."""
    if tid is None:
        return None
    e = nach_id(tid)
    if e is None:
        return "unsicher"
    return "unsicher" if e["gruppe"] == "grenzfall" else e["gruppe"]


def lesbar(tid: Optional[str]) -> Optional[str]:
    """Bezeichnung für Faktenbestätigung und Antwort (System-Text, nie vom Modell)."""
    if tid is None:
        return None
    if tid == UNBEKANNT:
        return "nicht im Regelwerk eingeordnet"
    if tid == WIDERSPRUCH:
        return "nicht eindeutig zuzuordnen"
    e = nach_id(tid)
    if e is None:
        return "nicht im Regelwerk eingeordnet"
    gruppe_text = {"kuenstlerisch": "künstlerisch/publizistisch", "nicht_kuenstlerisch": "nicht künstlerisch",
                   "grenzfall": "Grenzfall"}[e["gruppe"]]
    return f"{e['name']} ({gruppe_text})"


def fuer_modell() -> str:
    """Liste für die Modellanweisung: nur ID und Bezeichnung, keine Gruppe -- das Modell soll
    die Beschreibung zuordnen, nicht über die Abgabepflicht entscheiden."""
    return "\n".join(f"{e['id']}: {e['name']}" for e in alle())


def beim_selben_kuenstler(tid: Optional[str]) -> tuple[str, dict[str, Any] | None]:
    """Wie zählt die Leistung, wenn DERSELBE Künstler sie neben einer künstlerischen Leistung
    abrechnet? 'zaehlt', 'zaehlt_nicht' oder 'unsicher' (unbekannte Tätigkeit: unsicher)."""
    e = nach_id(tid)
    if e is None:
        return "unsicher", None
    return e["selber_kuenstler"], e
