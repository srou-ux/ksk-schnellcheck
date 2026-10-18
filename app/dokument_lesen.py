"""Lokaler Dokumentleser ohne Sprachmodell. Aufruf nur in einem begrenzten Arbeitsprozess.

Er liefert Text und unverbindliche Positionskandidaten, niemals eine fachliche Entscheidung.
PDF-Text: pypdf; Foto/Scan: lokales Tesseract/Poppler; Tabellen: openpyxl/csv.
Keine Makros/Formeln/Links ausführen. Originale liegen nur im temporären Arbeitsverzeichnis.
"""
from __future__ import annotations

import csv
import io
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import zipfile
from decimal import Decimal, InvalidOperation

KONFIG = json.loads(Path(__file__).with_name("dokument_config.json").read_text(encoding="utf-8"))


class Lesefehler(ValueError):
    pass


def geld(text):
    """Nur eindeutige Beträge, ohne Schätzung oder Rundung. Zahlen aus Excel sind Dezimalzahlen."""
    if isinstance(text, bool):
        return None
    if isinstance(text, (int, float, Decimal)):
        s = str(text)
    else:
        s = str(text).strip().replace("€", "").replace("EUR", "").strip()
        if re.fullmatch(r"-?(?:\d{1,3}(?:\.\d{3})+|\d+),\d{2}", s):
            s = s.replace(".", "").replace(",", ".")
        elif not re.fullmatch(r"-?\d+(?:\.\d{2})?", s):
            return None
    try:
        d = Decimal(s)
        if not d.is_finite() or d < 0 or d != d.quantize(Decimal("0.01")):
            return None
        return str(d)
    except InvalidOperation:
        return None


def textpositionen(text, beleg):
    kandidaten = []
    # Nur eine eindeutige Geldzahl am Zeilenende. Mengen-/Einzelpreisspalten nicht zusammenrechnen.
    zahl = r"-?(?:\d{1,3}(?:\.\d{3})+|\d+)[,.]\d{2}"
    for n, zeile in enumerate(text.splitlines(), start=1):
        if re.search(KONFIG["fremdwaehrung"], zeile, re.I):
            continue
        m = re.fullmatch(r"\s*(.+?)\s+(" + zahl + r")\s*(?:€|EUR)?\s*", zeile)
        if not m or re.search(zahl, m[1]) or re.search(KONFIG["summenzeile"], m[1], re.I):
            continue
        betrag = geld(m[2])
        if betrag is not None:
            bezeichnung = re.sub(r"^\d+\s{2,}", "", m[1].strip())
            kandidaten.append({"bezeichnung": bezeichnung, "betrag": betrag,
                               "beleg": f"{beleg}, Zeile {n}: {zeile.strip()}"})
    return kandidaten


def formularvorschlaege(text):
    """Nur beschriftete Dokumentfelder. Kein Rechnungsdatum als Zahlungsjahr interpretieren."""
    gefunden = {}
    zeilen = text.splitlines()
    labels = {"rechnungsempfänger", "rechnungsempfaenger", "auftraggeber"}
    for i, zeile in enumerate(zeilen):
        spalten = re.split(r"\s{2,}", zeile.strip())
        for spalte, zelle in enumerate(spalten):
            wert = None
            label, trenner, inhalt = zelle.partition(":")
            if label.strip().casefold() not in labels:
                continue
            if trenner and inhalt.strip():
                wert = inhalt.strip()
            elif i + 1 < len(zeilen):
                # Nur die nächste nichtleere Zeile im selben Spaltenraster.
                folge = next((z for z in zeilen[i + 1:] if z.strip()), "")
                teile = re.split(r"\s{2,}", folge.strip())
                if len(teile) == len(spalten):
                    wert = teile[spalte].strip()
            kopfwoerter = labels | {"rechnungssteller", "rechnungsstellerin", "leistung", "netto", "betrag"}
            if (wert and len(wert) <= 300 and wert.casefold() not in kopfwoerter
                    and not re.search(r":|\b(?:EUR|IBAN)\b", wert, re.I)):
                gefunden.setdefault("unternehmensart", []).append({"wert": wert,
                    "beleg": f"{zelle.strip()} → {wert}"})
    # Widersprechende Empfänger, z.B. mehrere Rechnungen in einer PDF: nichts vorbelegen.
    return {k: vs[0] for k, vs in gefunden.items() if len({v["wert"] for v in vs}) == 1}


def _ocr(bild: Path):
    if not shutil.which("tesseract"):
        raise Lesefehler("Texterkennung fehlt auf dem Server (Tesseract). Bitte eine Text-PDF oder Excel-Datei verwenden.")
    lang = subprocess.run(["tesseract", "--list-langs"], capture_output=True, text=True,
                          timeout=KONFIG["ocr_zeitlimit_s"], check=True).stdout.splitlines()
    sprache = "+".join(x for x in KONFIG["ocr_sprache"].split("+") if x in lang)
    if not sprache:
        raise Lesefehler("Die benötigten Tesseract-Sprachdaten fehlen auf dem Server.")
    r = subprocess.run(["tesseract", str(bild), "stdout", "-l", sprache, "--psm", "6"],
                       capture_output=True, text=True, timeout=KONFIG["ocr_zeitlimit_s"], check=True,
                       env={**os.environ, "OMP_THREAD_LIMIT": "1"})
    return r.stdout


def _tabellen(rows, name):
    text, kandidaten, hinweise = [], [], []
    kopf = None
    zellen = 0
    for n, row in enumerate(rows, start=1):
        zellen += len(row)
        if zellen > KONFIG["max_zellen"]:
            raise Lesefehler("Die Tabelle ist zu groß. Bitte in kleinere Dateien aufteilen.")
        vals = ["" if c is None else str(c) for c in row]
        zeile = " | ".join(vals)
        text.append(zeile)
        if sum(len(x) for x in text) > KONFIG["max_text"]:
            raise Lesefehler("Die Tabelle enthält zu viel Text. Bitte aufteilen.")
        if any(v.lstrip().startswith("=") for v in vals):
            hinweise.append(f"{name}, Zeile {n}: Formel nicht übernommen; Betrag bitte anhand des Originals ergänzen.")
            continue
        if re.search(KONFIG["fremdwaehrung"], zeile, re.I):
            hinweise.append(f"{name}, Zeile {n}: Fremdwährung nicht umgerechnet; tatsächlich gezahlten Euro-Betrag bitte ergänzen.")
            continue
        norm = [v.strip().lower() for v in vals]
        if kopf is None:
            bs = [i for i, v in enumerate(norm) if v in KONFIG["beschreibung_spalten"]]
            gs = [i for i, v in enumerate(norm) if v in KONFIG["betrag_spalten"]]
            if len(bs) == len(gs) == 1:
                kopf = bs[0], gs[0]
            continue
        b, g = kopf
        if max(b, g) >= len(row) or not vals[b].strip():
            continue
        if re.search(KONFIG["summenzeile"], vals[b], re.I):
            continue
        betrag = geld(row[g])
        if betrag is None:
            hinweise.append(f"{name}, Zeile {n}: Betrag nicht eindeutig; Position bitte prüfen und ergänzen.")
            continue
        kandidaten.append({"bezeichnung": vals[b].strip(), "betrag": betrag,
                           "beleg": f"{name}, Zeile {n}: {zeile}"})
    if kopf is None:
        hinweise.append(f"{name}: Spalten nicht eindeutig erkannt. Bitte Positionen manuell ergänzen.")
    return "\n".join(text), kandidaten, hinweise


def lese(pfad: Path, endung: str):
    if endung not in KONFIG["endungen"] or pfad.stat().st_size > KONFIG["max_bytes"]:
        raise Lesefehler("Dateityp oder Dateigröße wird nicht unterstützt.")
    texte, kandidaten, hinweise = [], [], []
    methode = "Text"
    if endung == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(pfad)
        if reader.is_encrypted:
            raise Lesefehler("Passwortgeschützte PDF: bitte eine ungeschützte Kopie hochladen.")
        if len(reader.pages) > KONFIG["max_seiten"]:
            raise Lesefehler("Die PDF hat zu viele Seiten. Bitte aufteilen.")
        for n, page in enumerate(reader.pages, start=1):
            # Normale Extraktion trennt häufig Beschreibung und Betrag derselben Tabellenzeile.
            # Layout erhält ihre räumliche Zuordnung; keine verketteten Nachbarzeilen erraten.
            text = page.extract_text(extraction_mode="layout") or ""
            if not text.strip():
                if not shutil.which("pdftoppm"):
                    raise Lesefehler("Diese PDF enthält einen Scan. Auf dem Server fehlt Poppler für die Texterkennung.")
                ziel = pfad.parent / f"seite-{n}"
                subprocess.run(["pdftoppm", "-f", str(n), "-l", str(n), "-r", str(KONFIG["ocr_dpi"]),
                                "-singlefile", "-png", str(pfad), str(ziel)], capture_output=True,
                               timeout=KONFIG["ocr_zeitlimit_s"], check=True)
                text = _ocr(ziel.with_suffix(".png"))
                methode = "Text und OCR"
            texte.append(f"Seite {n}\n{text}")
            kandidaten.extend(textpositionen(text, f"Seite {n}"))
    elif endung in (".png", ".jpg", ".jpeg"):
        from PIL import Image, ImageOps
        Image.MAX_IMAGE_PIXELS = KONFIG["max_bild_pixel"]
        with Image.open(pfad) as bild:
            if bild.width * bild.height > KONFIG["max_bild_pixel"] or getattr(bild, "n_frames", 1) != 1:
                raise Lesefehler("Das Bild ist zu groß oder enthält mehrere Einzelbilder.")
            ziel = pfad.parent / "ocr.png"
            ImageOps.exif_transpose(bild).convert("RGB").save(ziel)
        text = _ocr(ziel)
        texte.append(text)
        kandidaten = textpositionen(text, "Bild")
        methode = "OCR"
    elif endung == ".xlsx":
        import openpyxl
        with zipfile.ZipFile(pfad) as z:
            if sum(x.file_size for x in z.infolist()) > KONFIG["max_zip_bytes"]:
                raise Lesefehler("Die entpackte Excel-Datei ist zu groß.")
            if any("vbaProject" in x.filename for x in z.infolist()):
                raise Lesefehler("Dateien mit Makros werden nicht unterstützt.")
        wb = openpyxl.load_workbook(pfad, read_only=True, data_only=False, keep_links=False)
        try:
            for ws in wb:
                if ws.max_row * ws.max_column > KONFIG["max_zellen"]:
                    raise Lesefehler("Das Tabellenblatt ist zu groß. Bitte aufteilen.")
                text, ks, hs = _tabellen(ws.values, ws.title)
                texte.append(f"Blatt {ws.title}\n{text}"); kandidaten.extend(ks); hinweise.extend(hs)
        finally:
            wb.close()
        methode = "Excel"
    else:
        text = pfad.read_text(encoding="utf-8-sig")
        try:
            dialekt = csv.Sniffer().sniff(text[:4096], delimiters=";,\t")
        except csv.Error:
            dialekt = csv.excel
        text, kandidaten, hinweise = _tabellen(csv.reader(io.StringIO(text), dialekt), "CSV")
        texte.append(text); methode = "CSV"
    text = "\n\n".join(texte)
    if len(text) > KONFIG["max_text"] or len(kandidaten) > KONFIG["max_positionen"]:
        raise Lesefehler("Das Dokument ist zu umfangreich. Bitte aufteilen.")
    if not text.strip():
        raise Lesefehler("Kein Text erkannt. Bitte ein schärferes Bild oder eine Text-PDF verwenden.")
    hinweise.insert(0, "Automatisch gelesene Angaben sind ein Entwurf. Bitte jede Position, fehlende Zeilen, Nettobeträge und Zahlung anhand des Originals prüfen.")
    if "OCR" in methode:
        hinweise.append("Texterkennung kann Ziffern und Wörter verwechseln. Alle Beträge am Original kontrollieren.")
    if not kandidaten:
        hinweise.append("Keine eindeutigen Positionen erkannt. Sie können die Positionen unten manuell erfassen.")
    return {"text": text, "positionen": kandidaten, "hinweise": hinweise, "methode": methode,
            "formularvorschlaege": formularvorschlaege(text)}


def main():
    # Ressourcenbegrenzung zusätzlich zum Wanduhr-Limit des Elternprozesses; unter Windows nur letzteres.
    try:
        import resource
        mem = KONFIG["speicher_mb"] * 1024 * 1024
        resource.setrlimit(resource.RLIMIT_AS, (mem, mem))
        resource.setrlimit(resource.RLIMIT_CPU, (KONFIG["zeitlimit_s"], KONFIG["zeitlimit_s"]))
    except ImportError:
        pass
    try:
        result = lese(Path(sys.argv[1]), sys.argv[2])
        print(json.dumps({"ok": True, **result}, ensure_ascii=False))
    except Lesefehler as exc:
        print(json.dumps({"ok": False, "fehler": str(exc)}, ensure_ascii=False))
    except Exception:
        print(json.dumps({"ok": False, "fehler": "Datei nicht lesbar oder Zeitlimit erreicht. Bitte Format und Inhalt prüfen."}))


if __name__ == "__main__":
    main()
