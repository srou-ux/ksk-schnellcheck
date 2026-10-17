"""Gemeinsame Gestaltung der Prüflisten für den Steuerberater (Kopfzeile, Eingabefelder, Druck).

Unverändert übernommen aus scripts/erzeuge_pruefliste_2.py (17.10.2026, Umbau „ohne KI“: der Erzeuger der Prüfliste 2
brauchte den Quellenindex des Chats und ist entfernt; scripts/erzeuge_pruefliste_2_nachtrag.py nutzt diese Gestaltung).
"""
from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

SCHRIFT = "Arial"


EINTRAG = ("Entscheidung", "Richtige Fassung", "Erwartete Fundstelle richtig?", "Weitere zulässige Fundstellen",
           "Richtiges gilt von", "Richtiges gilt bis", "Richtige Stufe", "Begründung / Änderung", "Antwort",
           "Geprüft von", "Datum")


_KOPF = PatternFill("solid", fgColor="1F3A5F")


_GELB = PatternFill("solid", fgColor="FFF2CC")


_RAND = Border(bottom=Side(style="thin", color="BFBFBF"))


def _blatt(wb, titel: str, spalten: list[tuple[str, int]]):
    ws = wb.create_sheet(titel)
    ws.append([s for s, _ in spalten])
    for i, (name, breite) in enumerate(spalten, 1):
        c = ws.cell(row=1, column=i)
        c.font = Font(name=SCHRIFT, bold=True, color="FFFFFF")
        c.fill = _KOPF
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(i)].width = breite
    ws.row_dimensions[1].height = 32
    return ws


def _fertig(ws, auswahl: dict[str, str]) -> None:
    kopf = [c.value for c in ws[1]]
    for zeile in ws.iter_rows(min_row=2):
        for c in zeile:
            c.font = Font(name=SCHRIFT, size=10)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = _RAND
            if kopf[c.column - 1] in EINTRAG:
                c.fill = _GELB
    for name, liste in auswahl.items():
        spalte = get_column_letter(kopf.index(name) + 1)
        dv = DataValidation(type="list", formula1=liste, allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{spalte}2:{spalte}{max(ws.max_row, 2)}")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    _druck(ws)


def _druck(ws) -> None:
    """Querformat, auf Seitenbreite; Kopfzeile auf jeder Seite (der Steuerberater druckt evtl. aus)."""
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"

