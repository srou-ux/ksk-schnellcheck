"""Erzeugt einen datierten NACHTRAG zur Prüfliste 2 (15.10.2026).

Die Prüfliste 2 (regelwerk/Pruefliste_2_Steuerberater_Entwurf.xlsx) ist schon beim Nutzer. Sie wird nicht ersetzt
(Arbeitsfassung 2 des Nutzers, Teil B B3: "Prüfliste nur dann ersetzen, wenn keine bereits übermittelte Fassung
überschrieben wird. Sonst datierten Nachtrag"). Der Nachtrag enthält nur die Fachfragen aus
aufgaben/pruefliste_2_fragen.json, die in der übermittelten Prüfliste noch NICHT stehen (die Nummern werden aus der
Datei gelesen, nicht im Code festgelegt). Ein schon vorhandener Nachtrag mit Eintragungen wird nie überschrieben.

Aufruf (Windows): python scripts/erzeuge_pruefliste_2_nachtrag.py
"""
from __future__ import annotations

import datetime
import importlib.util
import json
import pathlib

import openpyxl
from openpyxl.styles import Alignment, Font

WURZEL = pathlib.Path(__file__).resolve().parents[1]
_spec = importlib.util.spec_from_file_location("pruefliste_2", WURZEL / "scripts" / "erzeuge_pruefliste_2.py")
P2 = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(P2)

UEBERMITTELT = P2.ZIEL
FRAGEN = WURZEL / "aufgaben" / "pruefliste_2_fragen.json"


def ziel(nummer: int) -> pathlib.Path:
    return WURZEL / "regelwerk" / f"Pruefliste_2_Nachtrag_{nummer}_Entwurf.xlsx"


def schon_vergeben(pfade: list[pathlib.Path]) -> set[str]:
    """Nummern der Fachfragen, die in der übermittelten Prüfliste oder früheren Nachträgen schon stehen."""
    nummern: set[str] = set()
    for pfad in pfade:
        if not pfad.exists():
            continue
        wb = openpyxl.load_workbook(pfad, read_only=True)
        if "Offene Fachfragen" in wb.sheetnames:
            nummern |= {str(r[0]) for r in wb["Offene Fachfragen"].iter_rows(min_row=2, values_only=True) if r and r[0]}
    return nummern


def baue(fragen: list[dict], stand: str, nummer: int) -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    anl = wb.active
    anl.title = "Anleitung"
    anl.column_dimensions["A"].width = 110
    titel = anl.cell(row=1, column=1, value=f"Prüfliste 2 – Nachtrag {nummer} für den Steuerberater – ENTWURF, Stand {stand}")
    titel.font = Font(name=P2.SCHRIFT, bold=True, size=14)
    texte = [
        "Dieser Nachtrag ergänzt die Prüfliste 2 (Datei Pruefliste_2_Steuerberater_Entwurf.xlsx). Er ersetzt sie nicht; "
        "die Prüfliste 2 bleibt gültig.",
        "Er enthält nur Fachfragen, die seit der Übergabe der Prüfliste 2 neu sind. Bitte die GELB hinterlegten Spalten "
        "ausfüllen (Antwort, Geprüft von, Datum).",
        "Alles sind Vorschläge des Systems; nichts ist freigegeben. Antworten werden als NEUE Version übernommen; nichts "
        "wird überschrieben.",
        f"Fragen in diesem Nachtrag: {len(fragen)} ({fragen[0]['nr']} bis {fragen[-1]['nr']})." if fragen else "Keine neuen Fragen.",
        "Erzeugt mit python scripts/erzeuge_pruefliste_2_nachtrag.py aus aufgaben/pruefliste_2_fragen.json.",
    ]
    for i, t in enumerate(texte, start=3):
        c = anl.cell(row=i, column=1, value=t)
        c.font = Font(name=P2.SCHRIFT, size=10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        anl.row_dimensions[i].height = 30
    ws = P2._blatt(wb, "Offene Fachfragen", [
        ("Nr.", 7), ("Thema", 20), ("Frage", 60), ("Was das System heute tut", 50), ("Bezug", 30), ("Herkunft", 24),
        ("Antwort", 50), ("Geprüft von", 12), ("Datum", 11)])
    for q in fragen:
        ws.append([q["nr"], q["thema"], q["frage"], q["heute"], q["bezug"], q["herkunft"], "", "", ""])
    P2._fertig(ws, {})
    P2._druck(anl)
    return wb


def main() -> None:
    fruehere = [UEBERMITTELT]
    nummer = 1
    while ziel(nummer).exists():
        fruehere.append(ziel(nummer))
        nummer += 1
    vergeben = schon_vergeben(fruehere)
    fragen = [q for q in json.loads(FRAGEN.read_text(encoding="utf-8"))["fragen"] if q["nr"] not in vergeben]
    if not fragen:
        print("Keine neuen Fachfragen; kein Nachtrag nötig.")
        return
    wb = baue(fragen, datetime.date.today().strftime("%d.%m.%Y"), nummer)
    wb.save(ziel(nummer))
    print("geschrieben:", ziel(nummer).relative_to(WURZEL), f"({len(fragen)} Fragen)")


if __name__ == "__main__":
    main()
