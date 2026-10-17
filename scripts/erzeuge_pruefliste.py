"""Erzeugt regelwerk/Pruefliste_Steuerberater_Entwurf.xlsx aus Regelkatalog und Testfällen.
Aufruf aus dem Projektstamm: python scripts/erzeuge_pruefliste.py
SCHUTZ: Sobald in der bestehenden Prüfliste eine Entscheidung eingetragen ist, bricht das Skript ab (sonst gingen
Entscheidungen des Steuerberaters verloren). Dann neue/geänderte Regeln von Hand in die Prüfliste übernehmen."""
import sys, pathlib
_alt = pathlib.Path('regelwerk/Pruefliste_Steuerberater_Entwurf.xlsx')
if _alt.exists():
    import openpyxl as _o
    _ws = _o.load_workbook(_alt)['Prüfliste']
    for _r in _ws.iter_rows(min_row=2, values_only=True):
        if (_r[11] not in (None, '', 'offen')) or _r[12] or _r[13]:
            sys.exit('ABBRUCH: In der Prüfliste stehen schon Entscheidungen. Nicht überschreiben.')
import openpyxl,re
from openpyxl.styles import Font,Alignment,PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
wb=openpyxl.load_workbook('regelwerk/Regelkatalog_24_KSVG_Entwurf.xlsx');ws=wb['Regeln']
h=[c.value for c in ws[1]];ix={k:i for i,k in enumerate(h)}
tw=openpyxl.load_workbook('regelwerk/Testfaelle_BSG_Urteile.xlsx')['Testfälle']
th=[c.value for c in tw[1]];ti={k:i for i,k in enumerate(th)}
tests={}
for r in tw.iter_rows(min_row=2,values_only=True):
    txt=' '.join(str(x) for x in r if x)
    for m in set(re.findall(r'R-\d+',txt)): tests.setdefault(m,[]).append(r[0])
rows=[]
for r in ws.iter_rows(min_row=2,values_only=True):
    v=lambda k:(r[ix[k]] or '')
    ps=v('Prüfstatus'); vb=v('Vorbehalt / Hinweis'); rf=v('Rechtsfolge')
    if 'fachlich zu klären' in ps or 'Grenzfall' in ps: p=1;g='1 – Zuerst: offen / Grenzfall'
    elif 'Auslegung der KSK' in vb or 'Ableitung' in vb or 'Ableitung' in rf or 'Auslegung der KSK' in rf: p=2;g='2 – KSK-Auslegung oder Ableitung (kein Gesetz)'
    else: p=3;g='3 – Gesetz / Rechtsprechung mit Fundstelle'
    t=tests.get(r[0],[])
    rows.append((p,g,r[0],v('Gruppe'),v('Norm'),v('Voraussetzung'),v('Schwelle / Zeitraum'),v('Gilt für Fassung'),rf,v('Fundstelle'),vb,', '.join(sorted(t)[:12])+(' …' if len(t)>12 else '')))
rows.sort(key=lambda x:(x[0],int(x[2].split('-')[1])))
out=openpyxl.Workbook();s=out.active;s.title='Prüfliste'
hdr=['Priorität','Regel-ID','Gruppe','Norm','Voraussetzung','Schwelle / Zeitraum','Gilt für Fassung','Rechtsfolge','Fundstelle','Vorbehalt / Hinweis','Tests dazu','Entscheidung','Begründung / Änderung','Geprüft von','Datum']
s.append(hdr)
for x in rows: s.append([x[1],x[2],x[3],x[4],x[5],x[6],x[7],x[8],x[9],x[10],x[11],'offen','','',''])
for c in s[1]:
    c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='1F3A5F');c.alignment=Alignment(wrap_text=True,vertical='top')
for i,w in enumerate([22,9,16,28,40,24,18,50,36,50,16,16,40,16,12],1): s.column_dimensions[get_column_letter(i)].width=w
for row in s.iter_rows(min_row=2):
    for c in row: c.alignment=Alignment(wrap_text=True,vertical='top')
dv=DataValidation(type='list',formula1='"offen,freigeben,mit Änderung freigeben,ablehnen,Rückfrage an KSK/DRV"',allow_blank=False);s.add_data_validation(dv);dv.add(f'L2:L{s.max_row}')
s.freeze_panes='C2';s.auto_filter.ref=s.dimensions
from collections import Counter
cn=Counter(x[1] for x in rows)
l=out.create_sheet('Anleitung')
for t in ['Prüfliste für den Steuerberater / Fachanwalt – ENTWURF. Stand 09.10.2026. Alle Regeln sind Vorschläge; nichts ist freigegeben.',
 f'{len(rows)} Regeln in drei Gruppen: '+'; '.join(f'{k}: {v}' for k,v in sorted(cn.items())),
 'So prüfen: Zeile lesen, Fundstelle in der Quelle nachschlagen (Dateiname siehe Regelkatalog, Spalte Quelldatei), Spalte „Entscheidung“ wählen, bei Änderung Begründung eintragen, Name und Datum setzen.',
 'Gruppe 1 zuerst: Regeln mit offenem Punkt (Datum 01.01.2023 am amtlichen Original) und der Grenzfall R-236 (8 Stunden).',
 'Gruppe 2: KSK-Auslegung oder Ableitung. Hier ist die Frage: Folgt das System der Verwaltungspraxis, und ist die Ableitung vertretbar?',
 'Gruppe 3: Wortlaut von Gesetz oder Urteil. Hier ist die Frage: Ist die Regel richtig wiedergegeben, stimmen Jahr und Fassung?',
 'Spalte „Tests dazu“: erfundene Testfälle (Regelwerk/Testfaelle_BSG_Urteile.xlsx), die diese Regel nennen. Eine Regel ohne Test ist kein Fehler, aber ein Hinweis auf geringere Absicherung.',
 'Erst nach Ihrer Freigabe darf ein Statuswechsel in „fachlich freigegeben“ erfolgen (Statusmodell_Entwurf.md). Das System macht Vorschläge, der Steuerberater entscheidet.']:
    l.append([t])
l.column_dimensions['A'].width=150
for row in l.iter_rows():
    for c in row: c.alignment=Alignment(wrap_text=True,vertical='top')
out.save('regelwerk/Pruefliste_Steuerberater_Entwurf.xlsx');print(cn,sum(1 for x in rows if not x[11]))
