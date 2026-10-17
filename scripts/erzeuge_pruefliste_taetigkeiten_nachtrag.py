"""Erzeugt regelwerk/Pruefliste_Taetigkeiten_Nachtrag_1_Entwurf.xlsx aus aufgaben/audit/taetigkeitskatalog_200.json
(15.10.2026). Die erste Tätigkeits-Prüfliste (regelwerk/Pruefliste_Taetigkeiten_Entwurf.xlsx) bleibt unverändert.
Bricht ab, wenn der Nachtrag schon Eintragungen des Steuerberaters enthält (nie überschreiben).

Aufruf (Windows): python scripts/erzeuge_pruefliste_taetigkeiten_nachtrag.py
"""
from __future__ import annotations

import datetime
import json
import pathlib
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

W = pathlib.Path(__file__).resolve().parents[1]
QUELLE = W / "aufgaben" / "audit" / "taetigkeitskatalog_200.json"
ZIEL = W / "regelwerk" / "Pruefliste_Taetigkeiten_Nachtrag_1_Entwurf.xlsx"
GELB = PatternFill("solid", fgColor="FFF2CC")
KOPF = PatternFill("solid", fgColor="1F3A5F")
EINTRAG = ["Entscheidung", "Neue oder geänderte Zeile im Blatt", "Begründung / Fundstelle", "Geprüft von", "Datum"]
LESBAR = {"K": "künstlerisch/publizistisch (zählt)", "N": "nicht künstlerisch (zählt nicht)", "G": "Grenzfall"}
URTEIL = {"ABWEICHUNG": "1 Abweichung: Blatt sagt etwas anderes als der Entwurf", "LUECKE": "2 Lücke: nicht im Blatt",
          "TREFFER": "3 Treffer: Blatt und Entwurf gleich"}
SYSTEM = {"A": "abgabepflichtig", "N": "nicht abgabepflichtig", "G": "Grenzfall", "R": "Rückfrage", "X": "nicht abgedeckt"}


def _schutz() -> None:
    if not ZIEL.exists():
        return
    ws = openpyxl.load_workbook(ZIEL)["Tätigkeiten"]
    kopf = [c.value for c in ws[1]]
    spalten = [i for i, k in enumerate(kopf) if k in EINTRAG]
    for z in ws.iter_rows(min_row=2, values_only=True):
        if any(z[i] not in (None, "", "offen") for i in spalten):
            sys.exit("ABBRUCH: Der Nachtrag enthält schon Eintragungen des Steuerberaters. Nicht überschreiben.")


def main() -> None:
    _schutz()
    k = json.loads(QUELLE.read_text(encoding="utf-8"))
    wb = openpyxl.Workbook()
    anl = wb.active
    anl.title = "Anleitung"
    anl.column_dimensions["A"].width = 110
    texte = [
        f"Tätigkeiten – Nachtrag 1 zur Prüfliste Tätigkeiten – ENTWURF, Stand {datetime.date.today():%d.%m.%Y}",
        "Zweck: 200 Dienstleistungen in demselben Rahmen (typischer Verwerter, Einzelperson, 2.000 €). Geprüft wird nur, "
        "ob die Tätigkeit als künstlerisch/publizistisch zählt.",
        "Spalte 'Entwurf': Einordnung aus dem Testkatalog des Nutzers (KI-gestützt, NICHT geprüft). Spalte 'Blatt': was das "
        "Excel-Blatt 'Tätigkeitsabgrenzung' heute sagt. Spalte 'System': was der Chatbot im Demo-Modus daraus macht.",
        "Bitte zuerst die Zeilen mit '1 Abweichung' und '2 Lücke' entscheiden (sortiert). Nur die GELBEN Spalten ausfüllen.",
        "Neu im Blatt seit 15.10.2026 (Entwurf): T-70 Textkorrektur, T-71 Unterricht in Musik/darstellender/bildender Kunst "
        "(KSVG § 2 'lehrt' gegen Checkliste 'Coaching'), T-72 technische Zeichnung, T-73 Playlist, T-74 Fotogrammetrie, "
        "T-75 Sync-Agent – alle außer T-70 als Grenzfall, bis Sie entscheiden.",
        "Ihre Entscheidungen werden als NEUE Zeilen ins Blatt übernommen; nichts wird überschrieben.",
    ]
    for i, t in enumerate(texte, 1):
        c = anl.cell(row=i, column=1, value=t)
        c.font = Font(name="Arial", size=14 if i == 1 else 10, bold=i == 1)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws = wb.create_sheet("Tätigkeiten")
    kopf = ["Nr.", "Tätigkeit", "Entwurf", "Hinweis im Entwurf", "Blatt (T-Zeile)", "Blatt (Gruppe)", "Vergleich",
            "System (Demo)", *EINTRAG]
    breiten = [8, 45, 26, 40, 10, 18, 34, 18, 26, 30, 40, 14, 11]
    ws.append(kopf)
    for i, b in enumerate(breiten, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = b
        c = ws.cell(row=1, column=i)
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        c.fill = KOPF
        c.alignment = Alignment(wrap_text=True)
    zeilen = sorted(k["faelle"], key=lambda f: (URTEIL[f["blatt"]["urteil"]], f["id"]))
    for f in zeilen:
        b = f["blatt"]
        ws.append([f["id"][3:], f["taetigkeit"], LESBAR[f["soll_kategorie"]], f["sicherheit"].split("(Entwurf des Nutzers")[1].split(")")[0].lstrip(": ")
                   if ": " in f["sicherheit"].split(")")[0] else "", b["t_id"] or "–", b["gruppe"] or "nicht im Blatt",
                   URTEIL[b["urteil"]], SYSTEM.get(f["demo_ergebnis"], f["demo_ergebnis"]), "offen", "", "", "", ""])
    dv = DataValidation(type="list", formula1='"offen,Entwurf richtig,Blatt richtig,anders (siehe Begründung)"', allow_blank=True)
    ws.add_data_validation(dv)
    for r in range(2, ws.max_row + 1):
        dv.add(f"I{r}")
        for c in range(9, 14):
            ws.cell(row=r, column=c).fill = GELB
        for c in range(1, 14):
            ws.cell(row=r, column=c).alignment = Alignment(wrap_text=True, vertical="top")
            ws.cell(row=r, column=c).font = Font(name="Arial", size=9)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(ZIEL)
    print("geschrieben:", ZIEL.relative_to(W), f"({len(zeilen)} Tätigkeiten)")


if __name__ == "__main__":
    main()
