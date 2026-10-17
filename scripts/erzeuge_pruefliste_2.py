"""Erzeugt regelwerk/Pruefliste_2_Steuerberater_Entwurf.xlsx (Schritt 3 des Plans, 14.10.2026).

Alles wird aus den Daten gelesen, nichts doppelt gepflegt:
- 'Regeln ohne Fassung': Regeln, deren Spalte „Gilt für Fassung“ der Quellenabruf keiner Fassung
  des Registers zuordnen kann (dieselbe Auswertung wie app/rag/index.py) -- sie werden nie geliefert.
- 'Fundstellen je Testfall': die Fehlschläge der Abruf-Güte (app/rag/guete.py) mit dem, was der
  Abruf stattdessen geliefert hat -- der Steuerberater legt fest, welche Fundstellen zulässig sind.
- 'Quellen Gültigkeit': app/rag/quellen_metadaten.csv (automatisch vorbelegt).
- 'Textbausteine neu': Bausteine des Regelkatalogs, die noch in keiner Prüfliste stehen.
- 'Offene Fachfragen': aufgaben/pruefliste_2_fragen.json.

Aufruf aus dem Projektstamm (Windows und Linux gleich; der Quellenindex wird bei Bedarf gebaut,
etwa 1 Minute):
    python scripts/erzeuge_pruefliste_2.py
SCHUTZ: Steht in der bestehenden Prüfliste schon eine Eintragung, bricht das Skript ab --
Entscheidungen des Steuerberaters dürfen nicht verloren gehen. Neue Fassung: neuer Dateiname.
"""
from __future__ import annotations

import csv
import datetime
import json
import pathlib
import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

WURZEL = pathlib.Path(__file__).resolve().parent.parent
sys.path.insert(0, str(WURZEL))
ZIEL = WURZEL / "regelwerk" / "Pruefliste_2_Steuerberater_Entwurf.xlsx"
KATALOG = WURZEL / "regelwerk" / "Regelkatalog_24_KSVG_Entwurf.xlsx"
FRAGEN = WURZEL / "aufgaben" / "pruefliste_2_fragen.json"
METADATEN = WURZEL / "app" / "rag" / "quellen_metadaten.csv"
FRUEHERE_PRUEFLISTEN = [WURZEL / "regelwerk" / "Pruefliste_Taetigkeiten_Entwurf.xlsx"]

OFFEN = "offen"
SCHRIFT = "Arial"
# Spalten, in die der Steuerberater schreibt (gelb hinterlegt; das Skript prüft sie beim Schutz).
EINTRAG = ("Entscheidung", "Richtige Fassung", "Erwartete Fundstelle richtig?", "Weitere zulässige Fundstellen",
           "Richtiges gilt von", "Richtiges gilt bis", "Richtige Stufe", "Begründung / Änderung", "Antwort",
           "Geprüft von", "Datum")
ENTSCHEIDUNG = '"offen,freigeben,mit Änderung freigeben,ablehnen,Rückfrage an KSK/DRV"'
STUFE_LESBAR = {"1": "1 Amtlich (Gesetz, Verordnung, Urteil)", "2": "2 Verwaltung (Auslegung der KSK/DRV, kein Gesetz)",
                "3": "3 Sekundär (nur Hinweis, nie Grundlage)"}

_KOPF = PatternFill("solid", fgColor="1F3A5F")
_GELB = PatternFill("solid", fgColor="FFF2CC")
_RAND = Border(bottom=Side(style="thin", color="BFBFBF"))


def _schutz(ziel: pathlib.Path = ZIEL) -> None:
    if not ziel.exists():
        return
    wb = openpyxl.load_workbook(ziel)
    for ws in wb.worksheets:
        kopf = [c.value for c in ws[1]]
        spalten = [i for i, k in enumerate(kopf) if k in EINTRAG]
        for zeile in ws.iter_rows(min_row=2, values_only=True):
            for i in spalten:
                wert = zeile[i] if i < len(zeile) else None
                if wert not in (None, "", OFFEN):
                    sys.exit(f"ABBRUCH: In '{ws.title}' stehen schon Eintragungen des Steuerberaters. Nicht überschreiben "
                             "(neue Fassung der Prüfliste unter neuem Dateinamen anlegen).")


# --------------------------------------------------------------------------------------------
# Daten
# --------------------------------------------------------------------------------------------
def fassungen_lesbar(fassungen: dict[int, str]) -> str:
    """{2019: 'A', …, 2023: 'B', …} -> 'A: 2019–2022; B: 2023–2027' (aus dem Register)."""
    teile: list[str] = []
    for f in sorted(set(fassungen.values()), key=lambda x: min(j for j, v in fassungen.items() if v == x)):
        jahre = sorted(j for j, v in fassungen.items() if v == f)
        teile.append(f"{f}: {jahre[0]}–{jahre[-1]}" if len(jahre) > 1 else f"{f}: {jahre[0]}")
    return "; ".join(teile)


def regeln_ohne_fassung(katalog: pathlib.Path = KATALOG, fassungen: dict[int, str] | None = None) -> list[dict]:
    from app.rag.index import Basis, _jahre_aus_fassung

    fassungen = fassungen if fassungen is not None else Basis().fassungen()
    ws = openpyxl.load_workbook(katalog, read_only=True, data_only=True)["Regeln"]
    zeilen = list(ws.iter_rows(values_only=True))
    ix = {k: i for i, k in enumerate(zeilen[0])}
    erg = []
    for z in zeilen[1:]:
        if not z or not z[0]:
            continue
        gilt = str(z[ix["Gilt für Fassung"]] or "")
        if not _jahre_aus_fassung(gilt, fassungen)[3]:
            erg.append({k: str(z[ix[k]] or "") for k in ("Regel-ID", "Gruppe", "Norm", "Voraussetzung", "Rechtsfolge",
                                                          "Fundstelle", "Gilt für Fassung", "Vorbehalt / Hinweis")})
    return erg


def fehlschlaege() -> list[dict]:
    from app.rag import guete
    from app.rag.index import Basis, baue

    basis = Basis()
    if not pathlib.Path(basis.index).exists():
        baue(basis)
    erg = []
    for fall, status, _angaben, oben in guete.miss(basis)["faelle"]:
        if status == "fehlschlag":
            erg.append({"id": fall.id, "gruppe": fall.gruppe, "jahr": fall.jahr, "frage": fall.suchtext,
                        "erwartet": fall.fundstelle, "geliefert": oben})
    return erg


def quellen(pfad: pathlib.Path = METADATEN) -> list[dict]:
    with open(pfad, encoding="utf-8", newline="") as d:
        return list(csv.DictReader(d))


def neue_textbausteine(katalog: pathlib.Path = KATALOG) -> list[tuple]:
    schon = set()
    for p in FRUEHERE_PRUEFLISTEN:
        if p.exists():
            wb = openpyxl.load_workbook(p, read_only=True)
            if "Textbausteine" in wb.sheetnames:
                schon |= {r[0] for r in wb["Textbausteine"].iter_rows(min_row=2, values_only=True) if r and r[0]}
    ws = openpyxl.load_workbook(katalog, read_only=True)["Textbausteine"]
    zeilen = list(ws.iter_rows(values_only=True))
    ix = {k: i for i, k in enumerate(zeilen[0])}
    # Neu sind die Bausteine mit höherer Nummer als alle, die schon in einer Prüfliste stehen (die
    # älteren Bausteine stammen aus der Zeit vor der ersten Prüfliste).
    nummern_schon = [int(str(t).split("-")[1]) for t in schon]
    ab = max(nummern_schon) + 1 if nummern_schon else 1
    return [(z[0], z[ix["Nur bei Ergebnisklasse (leer = immer)"]], z[ix["Text (mit {PLATZHALTERN})"]], z[ix["Bemerkung"]],
             z[ix["Prüfstatus"]])
            for z in zeilen[1:] if z and z[0] and int(str(z[0]).split("-")[1]) >= ab]


# --------------------------------------------------------------------------------------------
# Excel
# --------------------------------------------------------------------------------------------
def _blatt(wb, titel: str, spalten: list[tuple[str, int]]):
    ws = wb.create_sheet(titel)
    ws.append([s for s, _ in spalten])
    for i, (name, breite) in enumerate(spalten, 1):
        c = ws.cell(row=1, column=i)
        c.font = Font(name=SCHRIFT, bold=True, color="FFFFFF")
        c.fill = _KOPF
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(i)].width = breite
    ws.row_dimensions[1].height = 32
    return ws


def _fertig(ws, auswahl: dict[str, str]) -> None:
    kopf = [c.value for c in ws[1]]
    for zeile in ws.iter_rows(min_row=2):
        for c in zeile:
            c.font = Font(name=SCHRIFT, size=10)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = _RAND
            if kopf[c.column - 1] in EINTRAG:
                c.fill = _GELB
    for name, liste in auswahl.items():
        spalte = get_column_letter(kopf.index(name) + 1)
        dv = DataValidation(type="list", formula1=liste, allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{spalte}2:{spalte}{max(ws.max_row, 2)}")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    _druck(ws)


def _druck(ws) -> None:
    """Querformat, auf Seitenbreite; Kopfzeile auf jeder Seite (der Steuerberater druckt evtl. aus)."""
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"


def baue(regeln: list[dict], faelle: list[dict], qu: list[dict], bausteine: list[tuple], fragen: list[dict],
         fassungen_text: str, stand: str) -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    anl = wb.active
    anl.title = "Anleitung"

    # --- Regeln ohne Fassung ----------------------------------------------------------------
    ws = _blatt(wb, "Regeln ohne Fassung", [
        ("Regel-ID", 9), ("Gruppe", 16), ("Norm", 24), ("Voraussetzung", 42), ("Rechtsfolge", 42), ("Fundstelle", 32),
        ("Heutiger Eintrag „Gilt für Fassung“", 26), ("Vorbehalt / Hinweis", 34), ("Fassungen laut Register", 20),
        ("Folge heute", 30), ("Entscheidung", 16), ("Richtige Fassung", 16), ("Begründung / Änderung", 34),
        ("Geprüft von", 12), ("Datum", 11)])
    for r in regeln:
        ws.append([r["Regel-ID"], r["Gruppe"], r["Norm"], r["Voraussetzung"], r["Rechtsfolge"], r["Fundstelle"],
                   r["Gilt für Fassung"], r["Vorbehalt / Hinweis"], fassungen_text,
                   "Der Quellenabruf liefert diese Regel nie (Fassung nicht zuzuordnen). Im Rechenkern unverändert.",
                   OFFEN, "", "", "", ""])
    _fertig(ws, {"Entscheidung": ENTSCHEIDUNG,
                 "Richtige Fassung": '"A und B,nur A,nur B,nur ältere Fassung (vor A),nicht mehr anwenden"'})

    # --- Fundstellen je Testfall -------------------------------------------------------------
    ws = _blatt(wb, "Fundstellen je Testfall", [
        ("Testfall", 9), ("Gruppe", 16), ("Jahr", 7), ("Frage (Suchtext)", 50), ("Erwartete Fundstelle (Testfalldatei)", 36),
        ("Vom Abruf geliefert (Plätze 1–5)", 60), ("Erwartete Fundstelle richtig?", 16),
        ("Weitere zulässige Fundstellen", 34), ("Begründung / Änderung", 30), ("Geprüft von", 12), ("Datum", 11)])
    for f in faelle:
        geliefert = "\n".join(f"{i}) {t}" for i, t in enumerate(f["geliefert"], 1)) or "(nichts geliefert: „Nicht abgedeckt“)"
        ws.append([f["id"], f["gruppe"], f["jahr"] if f["jahr"] is not None else "–", f["frage"], f["erwartet"],
                   geliefert, OFFEN, "", "", "", ""])
    _fertig(ws, {"Erwartete Fundstelle richtig?": '"offen,ja,ja und weitere zulässig,teilweise,nein"'})

    # --- Quellen Gültigkeit ------------------------------------------------------------------
    ws = _blatt(wb, "Quellen Gültigkeit", [
        ("Datei", 40), ("Kurzname (für die Fundstelle)", 26), ("Art", 14), ("Stufe (Vorschlag)", 24),
        ("gilt von (Jahr)", 10), ("gilt bis (Jahr)", 10), ("aufnehmen", 10), ("nachrangig", 10), ("Bemerkung", 40),
        ("Entscheidung", 16), ("Richtiges gilt von", 11), ("Richtiges gilt bis", 11), ("Richtige Stufe", 12),
        ("Begründung / Änderung", 30), ("Geprüft von", 12), ("Datum", 11)])
    for q in qu:
        ws.append([q["datei"], q["kurzname"], q["art"], STUFE_LESBAR.get(q["stufe"], q["stufe"]),
                   q["gilt_von"] or "offen nach unten", q["gilt_bis"] or "bis heute", q["aufnehmen"], q["nachrangig"],
                   q["bemerkung"], OFFEN, "", "", "", "", "", ""])
    _fertig(ws, {"Entscheidung": ENTSCHEIDUNG, "Richtige Stufe": '"1,2,3,nicht verwenden"'})

    # --- Textbausteine neu -------------------------------------------------------------------
    ws = _blatt(wb, "Textbausteine neu", [
        ("Baustein-ID", 10), ("Erscheint bei", 24), ("Text (Platzhalter in {…} setzt das System)", 80), ("Bemerkung", 40),
        ("Prüfstatus im Regelkatalog", 22), ("Entscheidung", 16), ("Begründung / Änderung", 36), ("Geprüft von", 12),
        ("Datum", 11)])
    for tid, klasse, text, bem, status in bausteine:
        ws.append([tid, klasse or "jeder passenden Antwort", text, bem or "", status or "", OFFEN, "", "", ""])
    _fertig(ws, {"Entscheidung": ENTSCHEIDUNG})

    # --- Offene Fachfragen -------------------------------------------------------------------
    ws = _blatt(wb, "Offene Fachfragen", [
        ("Nr.", 7), ("Thema", 20), ("Frage", 60), ("Was das System heute tut", 50), ("Bezug", 30), ("Herkunft", 24),
        ("Antwort", 50), ("Geprüft von", 12), ("Datum", 11)])
    for q in fragen:
        ws.append([q["nr"], q["thema"], q["frage"], q["heute"], q["bezug"], q["herkunft"], "", "", ""])
    _fertig(ws, {})

    # --- Anleitung mit Übersicht (Formeln zählen, was noch offen ist) -------------------------
    anl.column_dimensions["A"].width = 34
    anl.column_dimensions["B"].width = 14
    anl.column_dimensions["C"].width = 14
    anl.column_dimensions["D"].width = 90
    titel = anl.cell(row=1, column=1, value=f"Prüfliste 2 für den Steuerberater – ENTWURF, Stand {stand}")
    titel.font = Font(name=SCHRIFT, bold=True, size=14)
    texte = [
        "Alles in dieser Datei sind Vorschläge des Systems; nichts ist freigegeben. Bitte nur die GELB hinterlegten Spalten ausfüllen "
        "(Auswahllisten, wo vorhanden). Name und Datum je Zeile setzen.",
        "Nach der Prüfung die Datei zurückgeben. Die Entscheidungen werden als NEUE Version in Regelkatalog, Metadaten und "
        "Testfälle übernommen; nichts wird überschrieben (alte Jahre werden weiter nach ihren Regeln gerechnet).",
        f"Fassungen laut Register: {fassungen_text} (Fassung A bis 31.12.2022, Fassung B ab 01.01.2023 – entscheidet die Excel).",
        "„Auslegung der KSK/DRV, kein Gesetz“: Verwaltungspraxis, nicht Gesetzestext. Rangfolge: Gesetz > Rechtsprechung > KSK/DRV-Auslegung.",
        "Beispiel für das Ausfüllen (Blatt „Fundstellen je Testfall“): Erwartete Fundstelle richtig? = „ja und weitere zulässig“; "
        "Weitere zulässige Fundstellen = „2) und 4)“ oder ein eigener Text wie „KSVG § 24 Abs. 1 Satz 1 Nr. 2“; "
        "Begründung = kurz, gern mit Fundstelle.",
        "Die erste Prüfliste (regelwerk/Pruefliste_Taetigkeiten_Entwurf.xlsx, Tätigkeiten) bleibt gültig und ist davon unabhängig. "
        "Weitere offene Punkte stehen im Regelkatalog, Blatt „Offene Punkte“.",
    ]
    for i, t in enumerate(texte, 3):
        c = anl.cell(row=i, column=1, value=t)
        anl.merge_cells(start_row=i, start_column=1, end_row=i, end_column=4)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = Font(name=SCHRIFT, size=10)
        anl.row_dimensions[i].height = 14 * (len(t) // 120 + 1)
    start = 3 + len(texte) + 1
    for j, k in enumerate(("Blatt", "Zeilen", "noch offen", "Was zu tun ist"), 1):
        c = anl.cell(row=start, column=j, value=k)
        c.font = Font(name=SCHRIFT, bold=True, color="FFFFFF")
        c.fill = _KOPF
    uebersicht = [
        ("Regeln ohne Fassung", "K", "Für jede Regel die Fassung festlegen (Spalte „Richtige Fassung“). Solange sie fehlt, liefert der Quellenabruf die Regel nie."),
        ("Fundstellen je Testfall", "G", "Der Abruf fand die erwartete Fundstelle nicht. Prüfen: Ist die erwartete Fundstelle richtig? Welche gelieferten Stellen sind ebenfalls zulässig?"),
        ("Quellen Gültigkeit", "J", "Jahre und Stufe jeder Quelle bestätigen oder berichtigen (automatisch vorbelegt). Bestimmt, welche Quelle für welches Jahr gilt."),
        ("Textbausteine neu", "F", "Diese Sätze erscheinen wörtlich in den Antworten. Fachlich richtig und verständlich?"),
        ("Offene Fachfragen", "G", "Fragen, die nur ein Berufsträger entscheiden kann. Kurz beantworten, gern mit Fundstelle."),
    ]
    for j, (blatt, spalte, was) in enumerate(uebersicht, start + 1):
        bez = f"'{blatt}'"
        werte = [blatt, f"=COUNTA({bez}!A:A)-1",
                 f'=COUNTIF({bez}!{spalte}:{spalte},"offen")+COUNTBLANK({bez}!{spalte}2:{spalte}{{ende}})' if blatt == "Offene Fachfragen"
                 else f'=COUNTIF({bez}!{spalte}:{spalte},"offen")', was]
        for k, v in enumerate(werte, 1):
            if isinstance(v, str) and "{ende}" in v:
                v = v.replace("{ende}", str(wb[blatt].max_row))
            c = anl.cell(row=j, column=k, value=v)
            c.font = Font(name=SCHRIFT, size=10)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    hinweis = anl.cell(row=start + len(uebersicht) + 2, column=1,
                       value="„noch offen“ zählt Zeilen mit „offen“ (bei den Fachfragen: ohne Antwort). Erzeugt mit "
                             "python scripts/erzeuge_pruefliste_2.py aus Regelkatalog, Quellenabruf, Metadaten und "
                             "aufgaben/pruefliste_2_fragen.json.")
    anl.merge_cells(start_row=hinweis.row, start_column=1, end_row=hinweis.row, end_column=4)
    hinweis.font = Font(name=SCHRIFT, size=9, italic=True)
    hinweis.alignment = Alignment(wrap_text=True)
    anl.row_dimensions[hinweis.row].height = 26
    for j in range(start + 1, start + 1 + len(uebersicht)):
        anl.row_dimensions[j].height = 28
    _druck(anl)
    return wb


def main() -> None:
    from app.rag.index import Basis

    _schutz()
    fassungen = Basis().fassungen()
    fragen = json.loads(FRAGEN.read_text(encoding="utf-8"))["fragen"]
    wb = baue(regeln_ohne_fassung(fassungen=fassungen), fehlschlaege(), quellen(), neue_textbausteine(), fragen,
              fassungen_lesbar(fassungen), datetime.date.today().strftime("%d.%m.%Y"))
    wb.save(ZIEL)
    print("geschrieben:", ZIEL.relative_to(WURZEL))


if __name__ == "__main__":
    main()
