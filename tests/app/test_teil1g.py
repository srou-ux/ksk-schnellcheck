"""Chatbot Teil 1g (Mischform: feste Textbausteine statt freiem Modelltext). Uebernimmt
jede Zeile aus `aufgaben/audit/abnahme_teil1g.py` als Unittest (G1-G7), plus zusaetzliche
Tests laut Auftrag Abschnitt D: Jahresgrenze eines Bausteins, fehlende/doppelte ID beim
Laden, Bausteinauswahl-Regeln einzeln, eine Fallbeschreibung mit Link/Paragraf/Ziffer/
Ergebnisaussage/HTML wird verworfen (die Antwort bleibt vollstaendig), Modellausfall bei
der Beschreibung, und ein Mischfall mit 4 Positionen (4.300,00 € / 215,00 €)."""
import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent.parent))

from kern import rechnen as K

from app import textbausteine
from app.baustein import erzeuge_platzhalter, pruefe_rohtext
from app.llm.base import AntwortKontext, LLMFehlgeschlagen
from app.llm.fake import FakeLLM
from app.schema import Fakten, Position as P
from app.textbausteine import Auswahl, Baustein, waehle_bausteine
from tests.app.hilfen import lauf_bestaetigt

NP, GMBH = "natuerliche_person", "GmbH_UG_AG"
WURZEL = Path(__file__).resolve().parent.parent.parent


def eur(x) -> str:
    return f"{D(x):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


class KeinKoerper(FakeLLM):
    """Wie in aufgaben/audit/abnahme_teil1g.py: formuliere_antwort darf nicht mehr existieren/
    aufgerufen werden, beschreibe_fall liefert einen vorgegebenen (ggf. absichtlich
    verbotenen) Text."""

    def __init__(self, beschreibung: str) -> None:
        super().__init__()
        self.beschreibung = beschreibung
        self.beschreibe_aufrufe = 0

    def beschreibe_fall(self, kontext):
        self.beschreibe_aufrufe += 1
        return self.beschreibung


GEMISCHT = Fakten(
    anfrageart="bewertung", jahr=2025, verwerterart="typischer_verwerter", leistungsart="Grafik/Design",
    positionen=[
        P(betrag=D("2400"), rolle="honorar", auftragnehmer_typ=NP),
        P(betrag=D("1100"), rolle="honorar", auftragnehmer_typ=NP),
        P(betrag=D("5000"), rolle="honorar", auftragnehmer_typ=GMBH, bezeichnung="GmbH-Studio"),
        P(betrag=D("800"), rolle="honorar", auftragnehmer_typ=NP),
    ],
)
BG_GEMISCHT = D("4300")


def _lauf(beschreibung: str, fakten: Fakten = None, frage: str = "F"):
    llm = KeinKoerper(beschreibung)
    ergebnis = lauf_bestaetigt(llm, frage, fakten or GEMISCHT.model_copy())
    return ergebnis, llm


class TestG1RegelkatalogBlatt(unittest.TestCase):
    """abnahme_teil1g.py G1/G1b."""

    def setUp(self):
        wb = openpyxl.load_workbook(
            WURZEL / "regelwerk" / "Regelkatalog_24_KSVG_Entwurf.xlsx", data_only=True,
        )
        self.wb = wb

    def test_blatt_textbausteine_vorhanden(self):
        self.assertIn("Textbausteine", self.wb.sheetnames)

    def test_jede_zeile_hat_id_text_und_pruefstatus(self):
        ws = self.wb["Textbausteine"]
        zeilen = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]
        self.assertTrue(zeilen)
        for zeile in zeilen:
            with self.subTest(id=zeile[0]):
                self.assertTrue(zeile[0])
                self.assertTrue(zeile[5])
                self.assertTrue(zeile[8])


class TestG2KeinFesttextImCode(unittest.TestCase):
    """abnahme_teil1g.py G2: kein Bausteintext (mit Platzhalter, laenger als 25 Zeichen)
    steht noch woertlich im Programmcode."""

    def test_keine_bausteintexte_im_app_code(self):
        wb = openpyxl.load_workbook(
            WURZEL / "regelwerk" / "Regelkatalog_24_KSVG_Entwurf.xlsx", data_only=True,
        )
        ws = wb["Textbausteine"]
        texte = [str(r[5]) for r in ws.iter_rows(min_row=2, values_only=True) if r[0] and r[5]]
        code = " ".join(p.read_text(encoding="utf-8") for p in (WURZEL / "app").glob("**/*.py"))
        fest = [t for t in texte if len(t) > 25 and "{" in t and t in code]
        self.assertTrue(texte)
        self.assertEqual(fest, [])


class TestG3KoerperOhneModell(unittest.TestCase):
    """abnahme_teil1g.py G3: Koerper kommt ohne Modell zustande, Zahlen aus dem Kern."""

    def test_antwort_ist_abgabepflichtig_nicht_fehlgeschlagen(self):
        ergebnis, _ = _lauf("Es geht um Aufträge an Freiberufler.")
        self.assertNotIn("fehlgeschlagen", ergebnis.ergebnisklasse.lower())
        self.assertIn("Abgabepflichtig", ergebnis.ergebnisklasse)

    def test_bemessungsgrundlage_und_abgabe_aus_dem_kern_stehen_drin(self):
        ab = K.berechne_abgabe(BG_GEMISCHT, 2025)[0]
        ergebnis, _ = _lauf("Es geht um Aufträge an Freiberufler.")
        self.assertIn(eur(BG_GEMISCHT), ergebnis.antwort)
        self.assertIn(eur(ab), ergebnis.antwort)

    def test_formuliere_antwort_existiert_nicht_mehr_auf_fakellm(self):
        self.assertFalse(hasattr(FakeLLM(), "formuliere_antwort"))

    def test_fundstelle_und_schlusssatz_vorhanden(self):
        ergebnis, _ = _lauf("Es geht um Aufträge an Freiberufler.")
        self.assertTrue(ergebnis.fundstellen)
        self.assertIn("Entwurf – vom Steuerberater zu prüfen", ergebnis.antwort)


class TestG4FallbeschreibungWirdGeprueft(unittest.TestCase):
    """abnahme_teil1g.py G4: verbotener Modelltext fliegt raus, Antwort bleibt vollstaendig;
    harmlose Beschreibung mit Platzhalter wird uebernommen."""

    def test_boese_beschreibung_mit_paragraf_aktenzeichen_betrag_link_fliegt_raus(self):
        boes = (
            "Das ist ganz sicher nicht abgabepflichtig, siehe § 54 und BSG 12 R 3/20, "
            "zahlen Sie 99 Euro, mehr unter http://ksk.example."
        )
        ergebnis, llm = _lauf(boes)
        ab = K.berechne_abgabe(BG_GEMISCHT, 2025)[0]
        self.assertGreaterEqual(llm.beschreibe_aufrufe, 1)
        for verbotenes in ("ganz sicher", "12 R 3/20", "99 Euro", "http"):
            with self.subTest(verbotenes=verbotenes):
                self.assertNotIn(verbotenes, ergebnis.antwort)
        self.assertNotIn("fehlgeschlagen", ergebnis.ergebnisklasse.lower())
        self.assertIn(eur(ab), ergebnis.antwort)

    def test_harmlose_beschreibung_mit_platzhalter_wird_uebernommen(self):
        ergebnis, _ = _lauf("Es geht um Zahlungen an Gestalter im Jahr {JAHR}.")
        self.assertIn("Zahlungen an Gestalter im Jahr 2025", ergebnis.antwort)

    def test_ergebnisaussage_im_modelltext_fliegt_raus(self):
        ab = K.berechne_abgabe(BG_GEMISCHT, 2025)[0]
        ergebnis, _ = _lauf("Damit ist die Sache nicht abgabepflichtig.")
        self.assertNotIn("Damit ist die Sache", ergebnis.antwort)
        self.assertIn(eur(ab), ergebnis.antwort)


class TestG5SaeumnisAuskunftFahrtkosten(unittest.TestCase):
    """abnahme_teil1g.py G5."""

    def test_saeumniszuschlag_wert_aus_dem_kern(self):
        ab = K.saeumniszuschlag([D("2000")], 4, 2025).abgabe
        ergebnis, _ = _lauf(
            "Es geht um eine Anfrage.",
            Fakten(anfrageart="saeumniszuschlag", jahr=2025, rueckstand=D("2000"), saeumnis_monate=4),
        )
        self.assertIn(eur(ab), ergebnis.antwort)
        self.assertNotIn("Abgabesatz", ergebnis.antwort)

    def test_auskunft_abgabesatz_2024_kommt_durch(self):
        ergebnis, _ = _lauf(
            "Es geht um eine Anfrage.",
            Fakten(anfrageart="auskunft", jahr=2024, auskunft_thema="abgabesatz"),
            frage="Wie hoch ist der Abgabesatz im Jahr 2024?",
        )
        self.assertIn("Auskunft", ergebnis.ergebnisklasse)
        self.assertIn("5", ergebnis.antwort)

    def test_fahrtkosten_bemessungsgrundlage_ohne_fahrtkosten(self):
        ergebnis, _ = _lauf(
            "Es geht um eine Anfrage.",
            Fakten(
                anfrageart="bewertung", jahr=2024, verwerterart="typischer_verwerter", auftragnehmer_typ=NP,
                leistungsart="Fotografie",
                positionen=[
                    P(betrag=D("1500"), rolle="honorar"),
                    P(betrag=D("300"), rolle="fahrtkosten_nachgewiesen", nachgewiesen=True),
                ],
            ),
            # Seit 11.10.2026 gilt eine Sonderrolle nur mit Beleg im Nutzertext (app/rollen_belege.json).
            frage="Fotograf 1.500 € Honorar, dazu 300 € nachgewiesene Fahrtkosten",
        )
        self.assertIn(eur(1500), ergebnis.antwort)
        self.assertNotIn(eur(1800), ergebnis.antwort.replace("1.500,00", ""))


class TestG6ModellausfallBeiDerBeschreibung(unittest.TestCase):
    """abnahme_teil1g.py G6: wirft beschreibe_fall LLMFehlgeschlagen, bleibt die Antwort
    trotzdem komplett (keine 'Pruefung fehlgeschlagen')."""

    class Defekt(KeinKoerper):
        def beschreibe_fall(self, kontext):
            raise LLMFehlgeschlagen("Zeitüberschreitung")

    def test_antwort_bleibt_trotz_modellausfall_komplett(self):
        ab = K.berechne_abgabe(BG_GEMISCHT, 2025)[0]
        llm = self.Defekt("")
        ergebnis = lauf_bestaetigt(llm, "G6 Modell fällt aus", GEMISCHT.model_copy())
        self.assertNotIn("fehlgeschlagen", ergebnis.ergebnisklasse.lower())
        self.assertIn(eur(ab), ergebnis.antwort)


class TestG7Schnittstelle(unittest.TestCase):
    """abnahme_teil1g.py G7: LLM-Schnittstelle und Mistral-Adapter kennen beschreibe_fall."""

    def test_llm_protocol_hat_beschreibe_fall(self):
        from app.llm.base import LLM

        self.assertIn("beschreibe_fall", dir(LLM))

    def test_mistral_adapter_hat_beschreibe_fall(self):
        from app.llm.mistral import MistralLLM

        self.assertTrue(hasattr(MistralLLM, "beschreibe_fall"))


class TestLadenFehlerBeiFehlenderOderDoppelterId(unittest.TestCase):
    """Abschnitt D: ein kaputtes Textbausteine-Blatt (doppelte ID / leerer Text / fehlender
    Pruefstatus) fuehrt beim Laden zu einem klaren Fehler, nicht zu stillem Weiterarbeiten."""

    def _schreibe(self, zeilen: list[tuple]) -> Path:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        ws = wb.create_sheet("Textbausteine")
        ws.append(["Baustein-ID", "Reihenfolge", "Benötigt", "Nicht, wenn vorhanden",
                    "Nur bei Ergebnisklasse", "Text", "Gültig von", "Gültig bis", "Prüfstatus", "Bemerkung"])
        for zeile in zeilen:
            ws.append(list(zeile))
        pfad = Path(tempfile.mkdtemp()) / "textbausteine_test.xlsx"
        wb.save(pfad)
        return pfad

    def test_doppelte_id_fuehrt_zu_fehler(self):
        pfad = self._schreibe([
            ("TB-X", 1, None, None, None, "Text eins.", None, None, "geprüft"),
            ("TB-X", 2, None, None, None, "Text zwei.", None, None, "geprüft"),
        ])
        with self.assertRaises(RuntimeError) as cm:
            textbausteine._lade(pfad)
        self.assertIn("doppelte ID", str(cm.exception))

    def test_leerer_text_fuehrt_zu_fehler(self):
        pfad = self._schreibe([("TB-X", 1, None, None, None, "   ", None, None, "geprüft")])
        with self.assertRaises(RuntimeError) as cm:
            textbausteine._lade(pfad)
        self.assertIn("Text ist leer", str(cm.exception))

    def test_fehlender_pruefstatus_fuehrt_zu_fehler(self):
        pfad = self._schreibe([("TB-X", 1, None, None, None, "Ein Text.", None, None, None)])
        with self.assertRaises(RuntimeError) as cm:
            textbausteine._lade(pfad)
        self.assertIn("Prüfstatus fehlt", str(cm.exception))

    def test_unbekannter_platzhalter_fuehrt_zu_fehler(self):
        pfad = self._schreibe([("TB-X", 1, "GEHEIMNISVOLL", None, None, "Ein Text.", None, None, "geprüft")])
        with self.assertRaises(RuntimeError) as cm:
            textbausteine._lade(pfad)
        self.assertIn("unbekannter Platzhalter", str(cm.exception))


class TestBausteinauswahlRegeln(unittest.TestCase):
    """Abschnitt D: die Auswahlregeln (benoetigt/nicht-wenn/ergebnisklasse/jahr) einzeln,
    mit injizierten Bausteinen (ueber den `bausteine=`-Parameter von waehle_bausteine) --
    keine Abhaengigkeit vom tatsaechlichen Regelkatalog-Inhalt."""

    def test_benoetigter_platzhalter_fehlt_baustein_passt_nicht(self):
        b = Baustein(id="X", reihenfolge=1, benoetigt=("ABGABE",), nicht_wenn=(), ergebnisklasse=None,
                     text="Abgabe: {ABGABE}.", gueltig_von=None, gueltig_bis=None)
        auswahl = waehle_bausteine({"JAHR": "2025"}, "Abgabepflichtig (Vorschlag)", 2025, bausteine=[b])
        self.assertTrue(auswahl.leerfall)

    def test_benoetigter_platzhalter_vorhanden_baustein_passt(self):
        b = Baustein(id="X", reihenfolge=1, benoetigt=("ABGABE",), nicht_wenn=(), ergebnisklasse=None,
                     text="Abgabe: {ABGABE}.", gueltig_von=None, gueltig_bis=None)
        auswahl = waehle_bausteine({"ABGABE": "73.50"}, "Abgabepflichtig (Vorschlag)", 2025, bausteine=[b])
        self.assertFalse(auswahl.leerfall)
        self.assertEqual(auswahl.texte, ["Abgabe: {ABGABE}."])

    def test_nicht_wenn_platzhalter_vorhanden_baustein_passt_nicht(self):
        b = Baustein(id="X", reihenfolge=1, benoetigt=(), nicht_wenn=("RUECKSTAND",), ergebnisklasse=None,
                     text="Kein Rueckstand.", gueltig_von=None, gueltig_bis=None)
        auswahl = waehle_bausteine({"RUECKSTAND": "2000"}, "x", None, bausteine=[b])
        self.assertTrue(auswahl.leerfall)

    def test_falsche_ergebnisklasse_baustein_passt_nicht(self):
        b = Baustein(id="X", reihenfolge=1, benoetigt=(), nicht_wenn=(), ergebnisklasse="Grenzfall",
                     text="Grenzfalltext.", gueltig_von=None, gueltig_bis=None)
        auswahl = waehle_bausteine({}, "Abgabepflichtig (Vorschlag)", 2025, bausteine=[b])
        self.assertTrue(auswahl.leerfall)

    def test_richtige_ergebnisklasse_baustein_passt(self):
        b = Baustein(id="X", reihenfolge=1, benoetigt=(), nicht_wenn=(), ergebnisklasse="Grenzfall",
                     text="Grenzfalltext.", gueltig_von=None, gueltig_bis=None)
        auswahl = waehle_bausteine({}, "Grenzfall", 2025, bausteine=[b])
        self.assertFalse(auswahl.leerfall)

    def test_gueltig_bis_vor_dem_anfragejahr_baustein_wird_nicht_verwendet(self):
        b = Baustein(id="X", reihenfolge=1, benoetigt=(), nicht_wenn=(), ergebnisklasse=None,
                     text="Alter Text.", gueltig_von=None, gueltig_bis=2022)
        auswahl = waehle_bausteine({}, "x", 2025, bausteine=[b])
        self.assertTrue(auswahl.leerfall)

    def test_gueltig_bis_nach_dem_anfragejahr_baustein_wird_verwendet(self):
        b = Baustein(id="X", reihenfolge=1, benoetigt=(), nicht_wenn=(), ergebnisklasse=None,
                     text="Alter Text.", gueltig_von=None, gueltig_bis=2030)
        auswahl = waehle_bausteine({}, "x", 2025, bausteine=[b])
        self.assertFalse(auswahl.leerfall)

    def test_gueltig_von_nach_dem_anfragejahr_baustein_wird_nicht_verwendet(self):
        b = Baustein(id="X", reihenfolge=1, benoetigt=(), nicht_wenn=(), ergebnisklasse=None,
                     text="Neuer Text.", gueltig_von=2030, gueltig_bis=None)
        auswahl = waehle_bausteine({}, "x", 2025, bausteine=[b])
        self.assertTrue(auswahl.leerfall)

    def test_ohne_anfragejahr_ein_jahresbegrenzter_baustein_passt_nicht(self):
        b = Baustein(id="X", reihenfolge=1, benoetigt=(), nicht_wenn=(), ergebnisklasse=None,
                     text="Text.", gueltig_von=2020, gueltig_bis=2030)
        auswahl = waehle_bausteine({}, "x", None, bausteine=[b])
        self.assertTrue(auswahl.leerfall)

    def test_leerfall_nimmt_den_bedingungslosen_baustein(self):
        bedingt = Baustein(id="X", reihenfolge=1, benoetigt=("NIE_VORHANDEN",), nicht_wenn=(), ergebnisklasse=None,
                            text="Wird nie verwendet.", gueltig_von=None, gueltig_bis=None)
        leer = Baustein(id="TB-LEER", reihenfolge=99, benoetigt=(), nicht_wenn=(), ergebnisklasse=None,
                         text="Dazu liegt kein passender Textbaustein vor.", gueltig_von=None, gueltig_bis=None)
        auswahl = waehle_bausteine({}, "x", 2025, bausteine=[bedingt, leer])
        self.assertTrue(auswahl.leerfall)
        self.assertEqual(auswahl.texte, [leer.text])

    def test_lade_sortiert_nach_reihenfolge_dann_id(self):
        # _lade() (nicht waehle_bausteine) ist fuer die Sortierung zustaendig -- hier mit
        # einem temporaeren Arbeitsblatt geprueft, in dem die Zeilen bewusst in der
        # "falschen" Reihenfolge stehen.
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        ws = wb.create_sheet("Textbausteine")
        ws.append(["Baustein-ID", "Reihenfolge", "Benötigt", "Nicht, wenn vorhanden",
                    "Nur bei Ergebnisklasse", "Text", "Gültig von", "Gültig bis", "Prüfstatus", "Bemerkung"])
        ws.append(["TB-B", 2, None, None, None, "Zweiter.", None, None, "geprüft", None])
        ws.append(["TB-A", 1, None, None, None, "Erster.", None, None, "geprüft", None])
        pfad = Path(tempfile.mkdtemp()) / "sortierung_test.xlsx"
        wb.save(pfad)
        geladen = textbausteine._lade(pfad)
        self.assertEqual([b.text for b in geladen], ["Erster.", "Zweiter."])

    def test_echter_regelkatalog_laedt_ohne_fehler_und_hat_genau_einen_leerfall_baustein(self):
        leerfaelle = [b for b in textbausteine.ALLE_BAUSTEINE if b.bedingungslos]
        self.assertEqual(len(leerfaelle), 1)


class TestFallbeschreibungMitVerbotenemInhaltWirdVerworfen(unittest.TestCase):
    """Abschnitt D: eine Fallbeschreibung mit Link, Paragraf, Ziffer, Ergebnisaussage oder
    HTML wird einzeln verworfen -- die Antwort selbst bleibt vollstaendig (kein
    'Pruefung fehlgeschlagen')."""

    FAELLE = {
        "link": "Mehr Informationen unter http://example.org/regel.",
        "paragraf": "Nach § 24 KSVG ist das so.",
        "ziffer": "Die Abgabe beträgt 999 Euro.",
        "ergebnisaussage": "Der Fall ist abgabepflichtig.",
        "html": "Wichtiger Hinweis: <b>bitte lesen</b>.",
    }

    def test_jeder_verbotene_inhalt_wird_verworfen_antwort_bleibt_vollstaendig(self):
        ab = K.berechne_abgabe(BG_GEMISCHT, 2025)[0]
        for name, zusatz in self.FAELLE.items():
            with self.subTest(fall=name):
                ergebnis, _ = _lauf(zusatz)
                self.assertNotIn(zusatz, ergebnis.antwort)
                self.assertNotIn("fehlgeschlagen", ergebnis.ergebnisklasse.lower())
                self.assertIn(eur(ab), ergebnis.antwort)


class TestModellfehlerBeiDerBeschreibung(unittest.TestCase):
    """Abschnitt D: jeder Fehler (nicht nur LLMFehlgeschlagen) beim Aufruf von
    beschreibe_fall lässt nur die Beschreibung entfallen."""

    class WirftUnerwarteteException(KeinKoerper):
        def beschreibe_fall(self, kontext):
            raise ValueError("unerwarteter Fehler (simuliert)")

    def test_unerwartete_exception_blockiert_die_antwort_nicht(self):
        ab = K.berechne_abgabe(BG_GEMISCHT, 2025)[0]
        llm = self.WirftUnerwarteteException("")
        ergebnis = lauf_bestaetigt(llm, "F", GEMISCHT.model_copy())
        self.assertNotIn("fehlgeschlagen", ergebnis.ergebnisklasse.lower())
        self.assertIn(eur(ab), ergebnis.antwort)


class TestMischfallVierPositionen(unittest.TestCase):
    """Abschnitt D: Mischfall mit 4 Positionen -- Bemessungsgrundlage 4.300,00 € (die
    5.000 €-GmbH-Position ist ausgeschlossen), Abgabe 215,00 € bei Abgabesatz 5 % (2023-2025)."""

    def test_bemessungsgrundlage_4300_und_abgabe_215(self):
        self.assertEqual(K.berechne_abgabe(BG_GEMISCHT, 2025)[0], D("215.00"))
        ergebnis, _ = _lauf("Es geht um Aufträge an Freiberufler.")
        self.assertIn("4.300,00", ergebnis.antwort)
        self.assertIn("215,00", ergebnis.antwort)


if __name__ == "__main__":
    unittest.main(verbosity=2)
