"""Führt die maschinenlesbaren Testfälle gegen den Regelkern aus und gleicht sie mit dem erwarteten Ergebnis in der Excel-Datei ab.
Aufruf:  python tests/lauf_testfaelle.py      (Exit-Code 1 bei Abweichung)
Entwurf – vom Steuerberater zu prüfen. Nicht ausgeführte Testfälle werden gezählt, nicht verschwiegen."""
import sys, json, pathlib, datetime
from decimal import Decimal as D
import openpyxl
W = pathlib.Path(__file__).resolve().parent.parent
sys.path.insert(0, str(W))
from kern import rechnen as K

KL = {"ABG": (K.ABGABEPFLICHTIG, "Abgabepflichtig"), "NICHT": (K.NICHT, "Nicht abgabepflichtig"),
      "GRENZ": (K.GRENZFALL, "Grenzfall"), "RUECK": (K.RUECKFRAGE, "Rückfrage")}


def de(x, cent=True):
    x = D(str(x)); s = f"{x:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return s


def varianten(x):
    x = D(str(x)); v = {de(x) + " €"}
    if x == x.to_integral_value():
        v |= {f"{int(x):,}".replace(",", ".") + " €", str(int(x)) + " €"}
    return v


def adapter(fn, a):
    jahr = a.get("jahr")
    if fn == "eigenwerbung": r = K.pruefe_eigenwerbung(jahr, a["summe"]); return r.klasse, r.bemessungsgrundlage, r.abgabe, None, None
    if fn == "typisch": r = K.pruefe_typischer_verwerter(jahr, a["betrag"]); return r.klasse, r.bemessungsgrundlage, r.abgabe, None, None
    if fn == "positionen": r = K.pruefe_positionen(jahr, a["positionen"]); return r.klasse, r.bemessungsgrundlage, r.abgabe, None, None
    if fn == "bemessung": r = K.bemessung_positionen(jahr, a["positionen"]); return r.klasse, r.bemessungsgrundlage, r.abgabe, None, None
    if fn == "veranstaltungen": r = K.veranstaltungen(jahr, a["anzahl"], a["summe"], a.get("einnahmeabsicht", True), a.get("art", "generalklausel")); return r.klasse, r.bemessungsgrundlage, r.abgabe, None, None
    if fn == "abgabe": x, _ = K.berechne_abgabe(a["bg"], jahr); return K.ABGABEPFLICHTIG, D(str(a["bg"])), x, None, None
    if fn == "vorauszahlung": x, r = K.vorauszahlung_monatlich(jahr, a["bg"]); return r.klasse, None, None, x, None
    if fn == "meldefrist": return None, None, None, None, K.meldefrist(a["jahr"])
    if fn == "aufbewahrung": return None, None, None, None, K.aufbewahrung_bis(a["jahr"])
    if fn == "saeumnis": r = K.saeumniszuschlag(a["rueckstand"], a["monate"], jahr, a.get("gesondert", False)); return r.klasse, None, r.abgabe, None, None
    if fn == "pauschale":
        bg, r = K.uebungsleiter_bg(jahr, a["betrag"], a["art"], True, a.get("bereits", 0)); return r.klasse, bg, r.abgabe, None, None
    if fn == "netto":
        n = K.netto(a["brutto"], a["ust"]); r = K.pruefe_typischer_verwerter(jahr, n); return r.klasse, r.bemessungsgrundlage, r.abgabe, None, None
    if fn == "kommission": b = K.kommission_bg(a["preis"], a["einbehalt"]); r = K.pruefe_typischer_verwerter(jahr, b); return r.klasse, b, r.abgabe, None, None
    if fn == "schaetzung": b = K.schaetzung_ausland(a["brutto"], jahr); r = K.pruefe_typischer_verwerter(jahr, b); return r.klasse, b, r.abgabe, None, None
    if fn == "rechnung":
        e, z = K.analysiere_rechnung(jahr, a["positionen"], a.get("empfaenger"), a.get("agentur"))
        return e.klasse, e.bemessungsgrundlage, e.abgabe, sum((x["offen"] for x in z), D(0)), None
    if fn == "verpflegung_bg":
        n, r = K.verpflegung(jahr, a["stunden"], False, a["erstattung"]); return r.klasse, D(str(a["honorar"])) + D(str(a["erstattung"])) - n, None, None, None
    raise KeyError(fn)


def main():
    spec = json.loads((W / "regelwerk" / "testfaelle_maschinenlesbar.json").read_text(encoding="utf-8"))
    ws = openpyxl.load_workbook(W / "regelwerk" / "Testfaelle_BSG_Urteile.xlsx")["Testfälle"]
    h = [c.value for c in ws[1]]; ix = {k: i for i, k in enumerate(h)}
    xl = {r[0].value: r for r in ws.iter_rows(min_row=2)}
    fehler = []; ok = 0
    for c in spec:
        i, s = c["id"], c["soll"]
        if i not in xl: fehler.append(f"{i}: nicht in der Excel-Datei"); continue
        text = str(xl[i][ix["Erwartetes Ergebnis"]].value or "")
        try:
            klasse, bg, abgabe, wert, datum = adapter(c["fn"], c["args"])
        except Exception as e:
            fehler.append(f"{i}: Kern-Fehler {e!r}"); continue
        probleme = []
        if "klasse" in s:
            if klasse != KL[s["klasse"]][0]: probleme.append(f"Kern liefert '{klasse}', Soll '{KL[s['klasse']][0]}'")
            if not text.startswith(KL[s["klasse"]][1]): probleme.append(f"Excel-Text beginnt nicht mit '{KL[s['klasse']][1]}': {text[:60]}")
        if "kern_klasse" in s and klasse != KL[s["kern_klasse"]][0]: probleme.append(f"Kern liefert '{klasse}', Soll '{KL[s['kern_klasse']][0]}'")
        if "text_enthaelt" in s and s["text_enthaelt"] not in text: probleme.append(f"Excel-Text enthält '{s['text_enthaelt']}' nicht")
        if "bg" in s:
            if bg != D(str(s["bg"])): probleme.append(f"Bemessungsgrundlage Kern {bg}, Soll {s['bg']}")
            if not any(v in text for v in varianten(s["bg"])): probleme.append(f"Excel-Text nennt {s['bg']} € nicht")
        if "abgabe" in s:
            if abgabe is None or abgabe != D(s["abgabe"]): probleme.append(f"Abgabe Kern {abgabe}, Soll {s['abgabe']}")
            if not any(v in text for v in varianten(s["abgabe"])): probleme.append(f"Excel-Text nennt Abgabe {s['abgabe']} € nicht")
        if "wert" in s:
            if wert is None or wert != D(s["wert"]): probleme.append(f"Wert Kern {wert}, Soll {s['wert']}")
            if not any(v in text for v in varianten(s["wert"])): probleme.append(f"Excel-Text nennt {s['wert']} € nicht")
        if "offen" in s and (wert is None or wert != D(str(s["offen"]))): probleme.append(f"Offener Betrag Kern {wert}, Soll {s['offen']}")
        if "offen" in s and not any(v in text for v in varianten(s["offen"])): probleme.append(f"Excel-Text nennt offen {s['offen']} € nicht")
        if "datum" in s:
            if datum.strftime("%d.%m.%Y") != s["datum"]: probleme.append(f"Datum Kern {datum}, Soll {s['datum']}")
            if s["datum"] not in text: probleme.append(f"Excel-Text nennt {s['datum']} nicht")
        if probleme: fehler.append(f"{i}: " + "; ".join(probleme))
        else: ok += 1
    gesamt = len(xl)
    print(f"{len(spec)} von {gesamt} Testfällen im Kern ausgeführt, {ok} bestanden, {len(fehler)} mit Abweichung; {gesamt - len(spec)} nur als Text (noch nicht im Kern).")
    for f in fehler: print("  ABWEICHUNG", f)
    return 1 if fehler else 0


if __name__ == "__main__":
    sys.exit(main())
